#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Sul de Minas weather by coffee season (Jul -> Jun) and ENSO phase, 2006/07 onward,
as an Excel workbook with native charts laid out like the "Temperature and
rainfall" El Niño / La Niña / Neutral panels.

Each variable comes from the source measured closest to the INMET stations
(tools/validate_sources.py -> validation/metrics.json). Region = mean of the
Sul de Minas towns in config.js.
    python tools/build_enso_workbook.py  ->  exports/SulDeMinas_ENSO_2006-2026.xlsx
"""
import json, os, sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
from fetch_reference import towns  # noqa: E402

V = os.path.join(ROOT, "validation")
OUT = os.path.join(ROOT, "exports", "SulDeMinas_ENSO_2006-2026.xlsx")
MONTHS = ["Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]
FIRST_SEASON = 2006

# variable -> (label, unit, number format, how to aggregate a month)
VARS = {
    "tmean": ("Mean temperature", "°C", "0.0", "mean"),
    "tmin": ("Min temperature", "°C", "0.0", "mean"),
    "tmax": ("Max temperature", "°C", "0.0", "mean"),
    "rain": ("Monthly rainfall", "mm", "0", "sum"),
    "rh": ("Relative humidity", "%", "0", "mean"),
    "soil": ("Soil moisture 0-100 cm", "% vol.", "0.0", "mean"),
    "soil_top": ("Surface soil moisture 0-7 cm", "% vol.", "0.0", "mean"),
}
SHEET = {"tmean": "Mean temp", "tmin": "Min temp", "tmax": "Max temp", "rain": "Rainfall",
         "rh": "Humidity", "soil": "Soil moisture"}
# Chosen source per variable — see validation/metrics.json and the Validation sheet
OMC = "Open-Meteo + INMET correction"
BEST = {"tmean": "Open-Meteo", "tmin": OMC, "tmax": OMC, "rain": "CHIRPS",
        "rh": "Open-Meteo", "soil": "Open-Meteo", "soil_top": "Open-Meteo"}
# label of the chosen source in the Validation sheet
CHOSEN = {"rain": "CHIRPS v3", "tmean": "Open-Meteo ERA5-Land", "tmin": "Open-Meteo ERA5-Land + INMET correction",
          "tmax": "Open-Meteo ERA5-Land + INMET correction", "rh": "Open-Meteo ERA5-Land"}
SOURCE_NAMES = {
    "CHIRPS": "CHIRPS v3.0 (UCSB Climate Hazards Center), 0.05° satellite + rain gauges",
    "Open-Meteo": "Open-Meteo, ERA5-Land 0.1° (models=era5_seamless), downscaled to town altitude",
    OMC: "Open-Meteo ERA5-Land, altitude-downscaled, + monthly offsets measured at 6 INMET stations "
         "(models smooth the daily extremes)",
    "NASA POWER": "NASA POWER (MERRA-2 / GEOS), 0.5° x 0.625°",
}

NAVY, GREY = "1F2A44", "D9D9D9"
HEAD = PatternFill("solid", fgColor=NAVY)
ALT = PatternFill("solid", fgColor="F2F4F8")
THIN = Border(bottom=Side(style="thin", color=GREY))
F = lambda **k: Font(name="Arial", **{"size": 10, **k})
PALETTE = ["1F77B4", "FF7F0E", "2CA02C", "E377C2", "7B3FA0", "222222", "17BECF", "BCBD22", "8C564B"]


# ------------------------------------------------------------------ data ----
def load(name):
    p = os.path.join(V, "grids", name + ".json")
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else None


def monthly(start, values, how):
    """Daily list -> {YYYY-MM: value}, complete months only."""
    by = defaultdict(list)
    d0 = date.fromisoformat(start)
    for i, v in enumerate(values):
        d = d0 + timedelta(i)
        by[d.strftime("%Y-%m")].append(v)
    out = {}
    for m, vals in by.items():
        y, mo = map(int, m.split("-"))
        n = ((date(y + mo // 12, mo % 12 + 1, 1)) - date(y, mo, 1)).days
        if len(vals) == n and all(v is not None for v in vals):
            out[m] = sum(vals) if how == "sum" else sum(vals) / n
    return out


def region_series():
    """{source: {var: {YYYY-MM: region mean}}} and per-town values."""
    ids = [t[0] for t in towns("sul_de_minas")]
    per_town = defaultdict(lambda: defaultdict(dict))        # source -> var -> town -> {month: v}
    om = load("om_towns")
    if om:
        for t in ids:
            s = om["series"][t]
            for var, key in (("tmean", "temperature_2m_mean"), ("tmin", "temperature_2m_min"),
                             ("tmax", "temperature_2m_max"), ("rain", "precipitation_sum"),
                             ("rh", "relative_humidity_2m")):
                per_town["Open-Meteo"][var][t] = monthly(om["start"], s[key], VARS[var][3])
            root = [None if None in (a, b, c) else 100 * (0.07 * a + 0.21 * b + 0.72 * c)
                    for a, b, c in zip(s["soil_moisture_0_to_7cm"], s["soil_moisture_7_to_28cm"],
                                       s["soil_moisture_28_to_100cm"])]
            per_town["Open-Meteo"]["soil"][t] = monthly(om["start"], root, "mean")
            top = [None if v is None else 100 * v for v in s["soil_moisture_0_to_7cm"]]
            per_town["Open-Meteo"]["soil_top"][t] = monthly(om["start"], top, "mean")
        corr = json.load(open(os.path.join(V, "metrics.json"), encoding="utf-8")).get("correction", {})
        for var, off in corr.items():                 # Tmin/Tmax: add station-measured monthly offsets
            for t in ids:
                per_town[OMC][var][t] = {m: v + off[str(int(m[5:]))]
                                         for m, v in per_town["Open-Meteo"][var][t].items()}
    nasa = load("nasa")
    if nasa:
        p = nasa["towns"]
        for t in ids:
            s = p["series"][t]
            for var, key, k in (("tmean", "T2M", 1), ("tmin", "T2M_MIN", 1), ("tmax", "T2M_MAX", 1),
                                ("rain", "PRECTOTCORR", 1), ("rh", "RH2M", 1),
                                ("soil", "GWETROOT", 100), ("soil_top", "GWETTOP", 100)):
                vals = [None if v is None else v * k for v in s[key]]
                per_town["NASA POWER"][var][t] = monthly(p["start"], vals, VARS[var][3])
    ch = load("chirps")
    if ch:
        for t in ids:
            per_town["CHIRPS"]["rain"][t] = {m: v for m, v in ch["series"][t].items() if v is not None}
    region = defaultdict(dict)
    for src, byvar in per_town.items():
        for var, bytown in byvar.items():
            months = set.intersection(*(set(m) for m in bytown.values()))
            region[src][var] = {m: sum(bytown[t][m] for t in ids) / len(ids) for m in months}
    return region, per_town, ids


def enso_phases():
    """Season start year -> (DJF ONI, phase). Rule: DJF >= +0.5 El Niño, <= -0.5 La Niña."""
    txt = load("nasa")["oni"]["text"]
    oni = {}
    for line in txt.splitlines()[1:]:
        p = line.split()
        if len(p) == 4:
            oni[(p[0], int(p[1]))] = float(p[3])
    last = list(oni.items())[-1]
    out = {}
    for y in range(FIRST_SEASON, date.today().year + 1):
        djf = oni.get(("DJF", y + 1))
        if djf is not None:
            out[y] = (djf, "El Niño" if djf >= 0.5 else "La Niña" if djf <= -0.5 else "Neutral", None)
        elif ("JJA", y) in oni or y == last[0][1]:
            v = last[1]
            ph = "El Niño*" if v >= 0.5 else "La Niña*" if v <= -0.5 else "Neutral*"
            out[y] = (None, ph, f"provisional: latest ONI {last[0][0]} {last[0][1]} = {v:+.2f}")
    return out, oni


def season_of(m):
    y, mo = map(int, m.split("-"))
    return y if mo >= 7 else y - 1


# ----------------------------------------------------------------- excel ----
def style_header(ws, row, ncol, height=30):
    for c in range(1, ncol + 1):
        cell = ws.cell(row, c)
        cell.font = F(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F(bold=True, size=14, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F(italic=True, size=9, color="555555")


def prepare():
    """Everything the workbook and the chart images need."""
    region, per_town, town_ids = region_series()
    phases, oni = enso_phases()
    best = {var: region[BEST[var]][var] for var in VARS}
    last_month = min(max(s) for s in best.values())          # last month every chosen source has
    months = []
    d = date(FIRST_SEASON, 7, 1)
    while d.strftime("%Y-%m") <= last_month:
        months.append(d.strftime("%Y-%m"))
        d = date(d.year + d.month // 12, d.month % 12 + 1, 1)
    seasons = sorted({season_of(m) for m in months})
    s_label = {y: f"{y}/{str(y + 1)[2:]}" for y in seasons}
    return region, phases, oni, best, last_month, months, seasons, s_label


def season_values(series, y):
    """12 values Jul..Jun of season y (None where missing)."""
    out = []
    for k in range(12):
        mo = (7 + k - 1) % 12 + 1
        out.append(series.get(f"{y + (1 if mo < 7 else 0)}-{mo:02d}"))
    return out


def build():
    region, phases, oni, best, last_month, months, seasons, s_label = prepare()
    metrics = json.load(open(os.path.join(V, "metrics.json"), encoding="utf-8"))

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    charts = wb.create_sheet("Charts")
    mon = wb.create_sheet("Monthly data")

    # ---- ENSO sheet (phase per season; other sheets link to it)
    ens = wb.create_sheet("ENSO")
    title(ens, "ENSO phase per coffee season (Jul-Jun)",
          "NOAA CPC Oceanic Niño Index (ONI). Rule: DJF ONI >= +0.5 El Niño, <= -0.5 La Niña, otherwise "
          "Neutral. * = season not finished, provisional.")
    hdr = ["Season", "DJF ONI", "ENSO phase", "SON ONI", "NDJ ONI", "Note"]
    for c, h in enumerate(hdr, 1):
        ens.cell(4, c, h)
    style_header(ens, 4, len(hdr))
    enso_row = {}
    for i, y in enumerate(seasons):
        r = 5 + i
        enso_row[y] = r
        djf, ph, note = phases.get(y, (None, "", None))
        ens.cell(r, 1, s_label[y])
        ens.cell(r, 2, djf).number_format = "+0.00;-0.00;0.00"
        ens.cell(r, 3, ph)
        ens.cell(r, 4, oni.get(("SON", y))).number_format = "+0.00;-0.00;0.00"
        ens.cell(r, 5, oni.get(("NDJ", y))).number_format = "+0.00;-0.00;0.00"
        ens.cell(r, 6, note or "")
    ens.cell(6 + len(seasons), 1, "Source: https://www.cpc.ncep.noaa.gov/data/indices/oni.ascii.txt").font = F(size=9, italic=True)
    for c, w in zip("ABCDEF", (10, 10, 13, 10, 10, 52)):
        ens.column_dimensions[c].width = w

    # ---- Monthly data (values; phase looked up in ENSO)
    title(mon, "Sul de Minas — monthly data (average of the towns), best source per variable",
          "Towns: " + ", ".join(t[1] for t in towns("sul_de_minas")) + ". Complete months only.")
    cols = ["Month", "Season", "ENSO phase"] + [f"{VARS[v][0]} ({VARS[v][1]})" for v in VARS]
    for c, h in enumerate(cols, 1):
        mon.cell(4, c, h)
    style_header(mon, 4, len(cols), 42)
    mon.cell(3, 4, "Source →").font = F(size=8, italic=True)
    for j, v in enumerate(VARS):
        mon.cell(3, 4 + j, BEST[v] if j else "Source → " + BEST[v]).font = F(size=8, italic=True, color="555555")
    mon_row = {}
    n_enso = len(seasons)
    for i, m in enumerate(months):
        r = 5 + i
        mon_row[m] = r
        y, mo = map(int, m.split("-"))
        mon.cell(r, 1, date(y, mo, 1)).number_format = "mmm yyyy"
        mon.cell(r, 2, s_label[season_of(m)])
        mon.cell(r, 3, f"=INDEX(ENSO!$C$5:$C${4 + n_enso},MATCH(B{r},ENSO!$A$5:$A${4 + n_enso},0))")
        for j, v in enumerate(VARS):
            val = best[v].get(m)
            c = mon.cell(r, 4 + j, None if val is None else round(val, 2))
            c.number_format = VARS[v][2]
        if i % 2:
            for c in range(1, len(cols) + 1):
                mon.cell(r, c).fill = ALT
    mon.freeze_panes = "B5"
    mon.column_dimensions["A"].width = 11
    mon.column_dimensions["B"].width = 9
    mon.column_dimensions["C"].width = 11
    for j in range(len(VARS)):
        mon.column_dimensions[get_column_letter(4 + j)].width = 14

    # ---- one sheet per variable: seasons x Jul..Jun, phase averages (formulas)
    var_sheet, avg_row = {}, {}
    col_of = {v: get_column_letter(4 + j) for j, v in enumerate(VARS)}
    for v, sheet_name in SHEET.items():
        ws = wb.create_sheet(sheet_name)
        var_sheet[v] = ws
        label, unit, fmt, how = VARS[v]
        title(ws, f"Sul de Minas — {label} ({unit}) by coffee season",
              f"Source: {SOURCE_NAMES[BEST[v]]}. Values link to 'Monthly data'.")
        hdr = ["Season", "ENSO phase"] + MONTHS + ["Season total" if how == "sum" else "Season mean"]
        for c, h in enumerate(hdr, 1):
            ws.cell(4, c, h)
        style_header(ws, 4, len(hdr))
        for i, y in enumerate(seasons):
            r = 5 + i
            ws.cell(r, 1, s_label[y]).font = F(bold=True)
            ws.cell(r, 2, f"=ENSO!C{enso_row[y]}")
            for k in range(12):
                mo = (7 + k - 1) % 12 + 1
                m = f"{y + (1 if mo < 7 else 0)}-{mo:02d}"
                if m in mon_row and best[v].get(m) is not None:
                    ws.cell(r, 3 + k, f"='Monthly data'!{col_of[v]}{mon_row[m]}").number_format = fmt
            agg = "SUM" if how == "sum" else "AVERAGE"
            ws.cell(r, 15, f'=IF(COUNT(C{r}:N{r})=12,{agg}(C{r}:N{r}),"")').number_format = fmt
        last = 4 + len(seasons)
        complete_last = 4 + sum(1 for y in seasons if f"{y + 1}-06" in mon_row)
        base = last + 2
        ws.cell(base, 1, "Averages").font = F(bold=True, color=NAVY)
        avg_row[v] = {}
        for k, ph in enumerate(["El Niño", "La Niña", "Neutral"]):
            r = base + 1 + k
            n = sum(1 for y in seasons if phases.get(y, (0, ""))[1] == ph)
            ws.cell(r, 1, f"{ph} average ({n} seasons)")
            for c in range(3, 16):
                L = get_column_letter(c)
                ws.cell(r, c, f'=AVERAGEIF($B$5:$B${last},"{ph}",{L}$5:{L}${last})').number_format = fmt
            avg_row[v][ph] = r
        r = base + 4
        ws.cell(r, 1, f"All complete seasons ({complete_last - 4})")
        for c in range(3, 16):
            L = get_column_letter(c)
            ws.cell(r, c, f"=AVERAGE({L}5:{L}{complete_last})").number_format = fmt
        for rr in range(base + 1, base + 5):
            for c in range(1, 16):
                ws.cell(rr, c).font = F(bold=True) if c == 1 else F(italic=True)
                ws.cell(rr, c).border = THIN
        ws.cell(base + 6, 1, "* season not finished — ENSO phase provisional, not included in the phase averages.").font = F(size=9, italic=True)
        ws.freeze_panes = "C5"
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 11
        for c in range(3, 16):
            ws.column_dimensions[get_column_letter(c)].width = 8.5
        ws.column_dimensions["O"].width = 12

    # ---- Charts: 3 phases x variables, like the reference slide
    title(charts, "Sul de Minas | Temperature, rainfall, humidity and soil moisture by ENSO phase",
          f"Coffee seasons Jul-Jun, {s_label[seasons[0]]} to {s_label[seasons[-1]]}. Dotted line = average "
          "of the phase. Best-validated source per variable (see Validation).")
    ranges = {}
    for v, ws in var_sheet.items():
        vals = [x for m, x in best[v].items() if m in mon_row]
        lo, hi = min(vals), max(vals)
        step = {"tmean": 2, "tmin": 2, "tmax": 2, "rain": 100, "rh": 10, "soil": 5}[v]
        ranges[v] = (step * int(lo // step), step * int(-(-hi // step)), step)
    order = ["tmean", "tmin", "tmax", "rain", "rh", "soil"]
    for i, v in enumerate(order):
        ws = var_sheet[v]
        cats = Reference(ws, min_col=3, max_col=14, min_row=4)
        for j, ph in enumerate(["El Niño", "La Niña", "Neutral"]):
            rows = [(y, 5 + seasons.index(y)) for y in seasons if phases.get(y, (0, ""))[1].rstrip("*") == ph]
            n_done = sum(1 for y, _ in rows if not phases[y][1].endswith("*"))
            ch = LineChart()
            ch.title = f"{ph} ({n_done} seasons) — {VARS[v][0]}"
            ch.y_axis.title = VARS[v][1]
            ch.height, ch.width = 8.2, 15.5
            ch.legend.position = "b"
            ch.display_blanks = "gap"
            lo, hi, step = ranges[v]
            ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = lo, hi, step
            ch.y_axis.number_format = VARS[v][2]
            ch.y_axis.majorGridlines.spPr = None
            ch.x_axis.delete = False
            ch.y_axis.delete = False
            for k, (y, r) in enumerate(rows):
                s = Series(Reference(ws, min_col=3, max_col=14, min_row=r), title=s_label[y] + (" *" if phases[y][1].endswith("*") else ""))
                s.graphicalProperties.line.solidFill = PALETTE[k % len(PALETTE)]
                s.graphicalProperties.line.width = 19050
                s.marker.symbol = "none"
                s.smooth = False
                ch.series.append(s)
            s = Series(Reference(ws, min_col=3, max_col=14, min_row=avg_row[v][ph]), title=f"{ph} average")
            s.graphicalProperties.line.solidFill = NAVY
            s.graphicalProperties.line.width = 31750
            s.graphicalProperties.line.dashStyle = "sysDot"
            s.marker.symbol = "none"
            s.smooth = False
            ch.series.append(s)
            ch.set_categories(cats)
            charts.add_chart(ch, f"{['A', 'J', 'S'][j]}{4 + i * 18}")
    for c in range(1, 28):
        charts.column_dimensions[get_column_letter(c)].width = 9.5

    # ---- INMET stations (observed; the truth used for validation)
    ins = wb.create_sheet("INMET stations")
    title(ins, "INMET automatic stations — observed monthly values (complete months only)",
          "portal.inmet.gov.br/dadoshistoricos (hourly archive, UTC days). Point measurements, not region averages.")
    st_ids = ["A515", "A531", "A529", "A524"]
    st = {c: json.load(open(os.path.join(V, "stations", c + ".json"), encoding="utf-8")) for c in st_ids}
    fields = [("rain", "Rain mm", "sum", "0"), ("tmean", "Tmean °C", "mean", "0.0"), ("tmin", "Tmin °C", "mean", "0.0"),
              ("tmax", "Tmax °C", "mean", "0.0"), ("rh", "RH %", "mean", "0")]
    ins.cell(4, 1, "Month")
    c = 2
    for sid in st_ids:
        ins.cell(3, c, f"{sid} {st[sid]['name'].title()} ({st[sid]['elev']:.0f} m)").font = F(bold=True, color=NAVY)
        for _, lab, _, _ in fields:
            ins.cell(4, c, lab)
            c += 1
    style_header(ins, 4, c - 1)
    st_month = {sid: {f: monthly(st[sid]["start"], st[sid]["series"][f], how) for f, _, how, _ in fields} for sid in st_ids}
    for i, m in enumerate(months):
        r = 5 + i
        y, mo = map(int, m.split("-"))
        ins.cell(r, 1, date(y, mo, 1)).number_format = "mmm yyyy"
        c = 2
        for sid in st_ids:
            for f, _, _, fmt in fields:
                val = st_month[sid][f].get(m)
                if val is not None:
                    ins.cell(r, c, round(val, 2)).number_format = fmt
                c += 1
    ins.freeze_panes = "B5"
    for cc in range(2, c):
        ins.column_dimensions[get_column_letter(cc)].width = 9

    # ---- Sources compared (region averages, every source)
    cmp_ = wb.create_sheet("Sources compared")
    title(cmp_, "Region monthly averages from every source (for comparison)",
          "Soil moisture: Open-Meteo = volumetric water (% vol.); NASA POWER = relative wetness GWETROOT/GWETTOP (%), not the same unit.")
    pairs = [(v, s) for v in VARS for s in ("CHIRPS", "Open-Meteo", OMC, "NASA POWER") if v in region.get(s, {})]
    cmp_.cell(4, 1, "Month")
    for k, (v, s) in enumerate(pairs):
        cmp_.cell(4, 2 + k, f"{VARS[v][0]} ({VARS[v][1]}) — {s}")
    style_header(cmp_, 4, 1 + len(pairs), 56)
    for i, m in enumerate(months):
        r = 5 + i
        y, mo = map(int, m.split("-"))
        cmp_.cell(r, 1, date(y, mo, 1)).number_format = "mmm yyyy"
        for k, (v, s) in enumerate(pairs):
            val = region[s][v].get(m)
            if val is not None:
                cmp_.cell(r, 2 + k, round(val, 2)).number_format = VARS[v][2]
    cmp_.freeze_panes = "B5"
    for k in range(len(pairs)):
        cmp_.column_dimensions[get_column_letter(2 + k)].width = 13

    # ---- Validation
    val = wb.create_sheet("Validation")
    title(val, "Which source is closest to reality? — comparison with INMET stations",
          metrics["period"] + ". 6 automatic stations in Minas Gerais coffee areas (Varginha, Maria da Fé, "
          "Passa Quatro, Formiga, Patrocínio, Manhuaçu), each source read at the station point.")
    hdr = ["Variable", "Source", "Months compared", "Bias (monthly)", "Mean abs. error (monthly)",
           "Error % (rain)", "Correlation r (monthly)", "Yearly error (rain, %)", "Daily MAE", "Daily r", "Chosen"]
    for c, h in enumerate(hdr, 1):
        val.cell(4, c, h)
    style_header(val, 4, len(hdr), 44)
    r = 5
    vnames = {"rain": "Rainfall (mm/month)", "tmean": "Mean temperature (°C)", "tmin": "Min temperature (°C)",
              "tmax": "Max temperature (°C)", "rh": "Relative humidity (%)"}
    for var in ("rain", "tmean", "tmin", "tmax", "rh"):
        res = metrics["by_var"][var]
        for src, rr in sorted(res.items(), key=lambda kv: (kv[1]["mg_monthly"] or {}).get("mae", 1e9)):
            m, y, dd = rr["mg_monthly"], rr["mg_yearly"], rr["mg_daily"]
            row = [vnames[var], src, m["n"], m["bias"], m["mae"], m.get("mae_pct"), m["r"],
                   y.get("mae_pct") if (y and var == "rain") else None,
                   dd["mae"] if dd else None, dd["r"] if dd else None]
            chosen = src == CHOSEN[var]
            row.append("✔" if chosen else "")
            for c, x in enumerate(row, 1):
                cell = val.cell(r, c, x)
                if c in (4, 5, 9):
                    cell.number_format = "0.00"
                if c in (7, 10):
                    cell.number_format = "0.000"
                if chosen:
                    cell.font = F(bold=True, color="1B5E20")
            r += 1
        r += 1
    rep = metrics.get("reports", {})
    if rep:
        val.cell(r, 1, "Second check — the monthly report tables of the site (towns of Sul de Minas and Cerrado)").font = F(bold=True, color=NAVY)
        r += 1
        for c, h in enumerate(["Variable", "Source", "Months compared", "Bias", "Mean abs. error", "Error %", "Correlation r"], 1):
            val.cell(r, c, h)
        style_header(val, r, 7, 30)
        r += 1
        for key, m in sorted(rep.items(), key=lambda kv: (kv[0].split("|")[1], kv[1]["mae"])):
            src, var = key.split("|")
            row = [vnames[var], src, m["n"], m["bias"], m["mae"], m.get("mae_pct"), m["r"]]
            for c, x in enumerate(row, 1):
                cell = val.cell(r, c, x)
                cell.number_format = "0.000" if c == 7 else "0.00" if c in (4, 5) else "General"
            r += 1
        r += 1
    notes = [
        "Bias = source − station (positive = source too high). Mean absolute error: lower is better. r: closer to 1 is better.",
        "Rainfall: CHIRPS blends satellite data with rain gauges (possibly including some of these stations), which is why it is closest.",
        "Temperatures: Open-Meteo is corrected to the real altitude of each point; NASA POWER uses a 50 km grid cell at its mean altitude.",
        "Min/max temperatures: reanalyses smooth the daily extremes (max too cold, min too warm). Monthly offsets measured at the "
        "stations are added (table below); the corrected row is tested leave-one-station-out (offsets computed without the tested station).",
        "Soil moisture: no station in the region measures it, so it cannot be verified here. ERA5-Land (Open-Meteo, 0.1°) is used: finer "
        "grid and forced by the same weather that validates best above; NASA POWER GWETROOT is in 'Sources compared'.",
        "Open-Meteo best_match (IFS 9 km blend, 2017+ only) is shown for information; the site/workbook use one consistent model (ERA5-Land).",
    ]
    for k, t in enumerate(notes):
        val.cell(r + 1 + k, 1, t).font = F(size=9, italic=True)
    corr = metrics.get("correction", {})
    if corr:
        r0 = r + 2 + len(notes)
        val.cell(r0, 1, "Monthly offsets added to Open-Meteo min/max temperatures (°C) — mean of station − model, "
                        "6 INMET stations, 2007-2026").font = F(bold=True, color=NAVY)
        for c, h in enumerate(["Month", "Tmin offset", "Tmax offset"], 1):
            val.cell(r0 + 1, c, h)
        style_header(val, r0 + 1, 3, 22)
        for mm in range(1, 13):
            val.cell(r0 + 1 + mm, 1, date(2000, mm, 1).strftime("%b"))
            for c, var in ((2, "tmin"), (3, "tmax")):
                if var in corr:
                    val.cell(r0 + 1 + mm, c, corr[var][str(mm)]).number_format = "+0.00;-0.00;0.00"
    for c, w in zip("ABCDEFGHIJK", (24, 24, 10, 11, 13, 10, 12, 12, 10, 9, 8)):
        val.column_dimensions[c].width = w

    # ---- README (French)
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    info = load("chirps")["info"] if load("chirps") else {}
    lines = [
        ("Sul de Minas — météo par saison caféière (juillet → juin) et phase ENSO", "h1"),
        (f"Fichier généré le {ts}. Données mensuelles de juillet {FIRST_SEASON} à {date(int(last_month[:4]), int(last_month[5:]), 1).strftime('%B %Y')} (mois complets uniquement).", ""),
        ("", ""),
        ("Zone", "h2"),
        ("Moyenne de 5 villes du Sul de Minas : " + ", ".join(t[1] for t in towns("sul_de_minas")) +
         " (mêmes coordonnées et altitudes que le site).", ""),
        ("", ""),
        ("Sources retenues (les plus proches des vraies mesures)", "h2"),
        ("Chaque source a été comparée aux stations automatiques INMET (mesures réelles, 2007–2026) : voir l'onglet Validation.", ""),
        ("• Pluie : CHIRPS v3.0 (UCSB) — satellite + pluviomètres, grille 0,05° (~5 km).", ""),
        ("• Température moyenne et humidité de l'air : Open-Meteo ERA5-Land (0,1°), corrigé à l'altitude de chaque ville.", ""),
        ("• Températures minimale et maximale : Open-Meteo ERA5-Land + correction mensuelle mesurée sur 6 stations INMET "
         "(les modèles lissent les extrêmes : maximales trop froides d'environ 2 °C, minimales trop chaudes d'environ 1 °C). "
         "Correction testée sur une station exclue du calcul à chaque fois (voir Validation).", ""),
        ("• Humidité du sol (0–100 cm et surface 0–7 cm) : Open-Meteo ERA5-Land, en % du volume de sol. "
         "Aucune station ne la mesure dans la région : c'est une estimation de modèle, non vérifiable localement.", ""),
        ("", ""),
        ("Phases ENSO", "h2"),
        ("Indice ONI de la NOAA pour décembre-janvier-février : ≥ +0,5 = El Niño, ≤ −0,5 = La Niña, sinon Neutre. "
         "Cette règle redonne exactement le classement du graphique de référence (2006/07–2024/25).", ""),
        ("2026/27 : saison en cours, El Niño en formation (ONI juin-juillet-août 2026 = +1,80) — marquée « * », provisoire, "
         "exclue des moyennes.", ""),
        ("", ""),
        ("Onglets", "h2"),
        ("Charts : graphiques par phase (comme le modèle), une ligne par saison, pointillés = moyenne de la phase.", ""),
        ("Monthly data : toutes les valeurs mensuelles. Mean temp … Soil moisture : tableaux saison × mois (formules liées à Monthly data) "
         "et moyennes par phase.", ""),
        ("INMET stations : mesures réelles de 4 stations (Varginha, Maria da Fé, Passa Quatro, Formiga). "
         "Sources compared : les 3 sources côte à côte. Validation : précision de chaque source. ENSO : indice par saison.", ""),
        ("", ""),
        ("Fraîcheur des données", "h2"),
        (f"CHIRPS final jusqu'à {info.get('final_through', '?')} ; préliminaire jusqu'à {info.get('prelim_through', '?')} (non utilisé ici). "
         "Open-Meteo ERA5 : ~5 jours de retard. NASA POWER : quelques jours.", ""),
        ("", ""),
        ("Crédits", "h2"),
        ("CHIRPS : Funk et al., Climate Hazards Center UCSB. Open-Meteo (CC BY 4.0) / ERA5-Land, Copernicus-ECMWF. "
         "NASA POWER, NASA Langley. INMET. NOAA CPC (ONI).", ""),
    ]
    for i, (t, kind) in enumerate(lines, 1):
        c = readme.cell(i, 1, t)
        c.font = F(bold=True, size=15, color=NAVY) if kind == "h1" else F(bold=True, size=11, color=NAVY) if kind == "h2" else F()
        c.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 130

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title not in ("README", "Charts")
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and (cell.font is None or cell.font.name != "Arial"):
                    cell.font = F(bold=cell.font.b, italic=cell.font.i, color=cell.font.color, size=cell.font.sz or 10)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT, "| months", months[0], "->", months[-1], "| seasons", len(seasons))
    return region, best, months, seasons, phases


if __name__ == "__main__":
    build()
