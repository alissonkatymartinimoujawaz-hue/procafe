"""Recria em Excel o grafico "Arabica x Robusta / ciclo bienal ON-OFF".

Gera excel/producao_arabica_robusta.xlsx: tabela de dados + grafico combinado
(colunas agrupadas + linha do ciclo bienal + tendencia da Robusta).
"""

import re
import zipfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.fill import PatternFillProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as DrawingFont,
    Paragraph,
    ParagraphProperties,
    RichTextProperties,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = "4472C4"        # Arabica
ORANGE = "ED7D31"      # Robusta
DARK = "1B2A41"        # linha do ciclo bienal
TREND = "E8A33D"       # tendencia (pontilhada) da Robusta
GRID = "D9D9D9"
AXIS = "BFBFBF"
TEXT = "404040"

# Valores digitalizados do grafico de origem (imagem fornecida pelo usuario),
# em mil sacas de 60 kg. Ciclo bienal alternando ON/OFF a partir de 2010/11.
ROWS = [
    ("2010/11", 37800, 11500, "ON"),
    ("2011/12", 35200, 14300, "OFF"),
    ("2012/13", 41000, 14800, "ON"),
    ("2013/14", 38600, 15200, "OFF"),
    ("2014/15", 37800, 16800, "ON"),
    ("2015/16", 36700, 13400, "OFF"),
    ("2016/17", 45400, 10200, "ON"),
    ("2017/18", 36400, 12600, "OFF"),
    ("2018/19", 49300, 18300, "ON"),
    ("2019/20", 38900, 17600, "OFF"),
    ("2020/21", 48800, 20100, "ON"),
    ("2021/22", 36800, 21000, "OFF"),
    ("2022/23", 39400, 22300, "ON"),
    ("2023/24", 45300, 21200, "OFF"),
    ("2024/25", 44400, 22100, "ON"),
    ("2025/26", 38300, 25400, "OFF"),
    ("2026/27", 49000, 24200, "ON"),
]

FIRST, LAST = 2, 1 + len(ROWS)

wb = Workbook()
ws = wb.active
ws.title = "Producao"

# ---------------------------------------------------------------- premissas
ws["J1"] = "Premissas do grafico (editaveis)"
ws["J1"].font = Font(name=FONT, size=10, bold=True, color=TEXT)
ws["J2"], ws["K2"] = "Nivel da linha em ano ON", 61000
ws["J3"], ws["K3"] = "Nivel da linha em ano OFF", 33000
ws["J4"], ws["K4"] = "Deslocamento do rotulo ON/OFF", 2200
for r in (2, 3, 4):
    ws[f"J{r}"].font = Font(name=FONT, size=10, color=TEXT)
    ws[f"K{r}"].font = Font(name=FONT, size=10, bold=True, color="0000FF")
    ws[f"K{r}"].fill = PatternFill("solid", fgColor="FFFF00")
    ws[f"K{r}"].number_format = "#,##0"

ws["J6"] = (
    "Serie 'Biennial ON/OFF pattern' = nivel ON/OFF conforme a coluna D. "
    "Serie 'Robusta progression' = reta de minimos quadrados (SLOPE/INTERCEPT) "
    "sobre a producao de Robusta."
)
ws["J6"].font = Font(name=FONT, size=9, italic=True, color=TEXT)
ws["J7"] = (
    "Fonte dos valores: digitalizados do grafico enviado pelo usuario "
    "(producao em mil sacas de 60 kg)."
)
ws["J7"].font = Font(name=FONT, size=9, italic=True, color=TEXT)

# ------------------------------------------------------------------ cabecalho
headers = [
    "Market Year",
    "Arabica Production",
    "Robusta Production",
    "Biennial cycle",
    "Biennial ON/OFF pattern",
    "Robusta progression",
    "Rotulo ON/OFF (posicao)",
    "Indice do ano",
]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="44546A")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 30

thin = Side(style="thin", color=GRID)
for i, (year, ara, rob, cycle) in enumerate(ROWS):
    r = FIRST + i
    ws.cell(row=r, column=1, value=year)
    ws.cell(row=r, column=2, value=ara)
    ws.cell(row=r, column=3, value=rob)
    ws.cell(row=r, column=4, value=cycle)
    ws.cell(row=r, column=5, value=f'=IF($D{r}="ON",$K$2,$K$3)')
    ws.cell(row=r, column=6, value=(
        f"=SLOPE($C${FIRST}:$C${LAST},$H${FIRST}:$H${LAST})*$H{r}"
        f"+INTERCEPT($C${FIRST}:$C${LAST},$H${FIRST}:$H${LAST})"
    ))
    ws.cell(row=r, column=7, value=f"=$B{r}+$K$4")
    ws.cell(row=r, column=8, value=i + 1)
    for col in range(1, 9):
        cell = ws.cell(row=r, column=col)
        cell.font = Font(name=FONT, size=10, color=TEXT)
        cell.border = Border(bottom=thin)
        if col in (2, 3, 5, 6, 7):
            cell.number_format = "#,##0"
        else:
            cell.alignment = Alignment(horizontal="center")

widths = {"A": 12, "B": 15, "C": 15, "D": 12, "E": 16, "F": 15, "G": 16, "H": 11,
          "J": 30, "K": 12}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# ------------------------------------------------------------------ graficos
def txpr(size=800, bold=False, color=TEXT):
    return RichText(
        bodyPr=RichTextProperties(),
        p=[Paragraph(
            pPr=ParagraphProperties(
                defRPr=CharacterProperties(
                    sz=size, b=bold, solidFill=color,
                    latin=DrawingFont(typeface=FONT),
                )
            ),
            endParaRPr=CharacterProperties(sz=size, b=bold, solidFill=color),
        )],
    )


# categorias como texto (strRef) para o Excel nao tentar ler "2010/11" como data
CATS_REF = f"'{ws.title}'!$A${FIRST}:$A${LAST}"

bar = BarChart()
bar.type = "col"
bar.grouping = "clustered"
bar.gapWidth = 55
bar.overlap = 0
for col, color in ((2, BLUE), (3, ORANGE)):
    ref = Reference(ws, min_col=col, min_row=1, max_row=LAST)
    s = Series(ref, title_from_data=True)
    s.graphicalProperties = GraphicalProperties(solidFill=color)
    s.graphicalProperties.line = LineProperties(noFill=True)
    s.cat = AxDataSource(strRef=StrRef(f=CATS_REF))
    bar.series.append(s)

line = LineChart()
# ciclo bienal (linha preta em zigue-zague)
s = Series(Reference(ws, min_col=5, min_row=1, max_row=LAST), title_from_data=True)
s.graphicalProperties = GraphicalProperties()
s.graphicalProperties.line = LineProperties(solidFill=DARK, w=19050)
s.marker = Marker(symbol="none")
s.smooth = False
line.series.append(s)

# tendencia da Robusta (linha laranja pontilhada)
s = Series(Reference(ws, min_col=6, min_row=1, max_row=LAST), title_from_data=True)
s.graphicalProperties = GraphicalProperties()
s.graphicalProperties.line = LineProperties(solidFill=TREND, w=15875, prstDash="dash")
s.marker = Marker(symbol="none")
s.smooth = False
line.series.append(s)

# serie invisivel que carrega os rotulos "ON" / "OFF" acima das barras Arabica
s = Series(Reference(ws, min_col=7, min_row=1, max_row=LAST), title_from_data=True)
s.graphicalProperties = GraphicalProperties()
s.graphicalProperties.line = LineProperties(noFill=True)
s.marker = Marker(symbol="none")
s.smooth = False
labels = []
for i, (_, _, _, cycle) in enumerate(ROWS):
    dl = DataLabel(idx=i)
    dl.numFmt = f'"{cycle}"'          # exibe o texto ON/OFF no lugar do valor
    dl.showVal = True
    dl.showSerName = dl.showCatName = dl.showLegendKey = dl.showPercent = False
    dl.showBubbleSize = False
    dl.dLblPos = "t"
    dl.txPr = txpr(size=650, bold=True, color=TEXT)
    labels.append(dl)
s.dLbls = DataLabelList(dLbl=labels, showVal=True, showSerName=False,
                        showCatName=False, showLegendKey=False,
                        showPercent=False, showBubbleSize=False)
line.series.append(s)
for s in line.series:
    s.cat = AxDataSource(strRef=StrRef(f=CATS_REF))

bar += line
chart = bar

chart.height = 10.5
chart.width = 26
chart.style = None
chart.title = None
chart.roundedCorners = False
chart.graphical_properties = GraphicalProperties(solidFill="FFFFFF")
chart.graphical_properties.line = LineProperties(noFill=True)

chart.x_axis.title = "Market Year"
chart.y_axis.title = "Production (1,000 bags)"
for ax in (chart.x_axis, chart.y_axis):
    ax.delete = False
    ax.txPr = txpr(size=750)
    ax.title.tx.rich.p[0].pPr = ParagraphProperties(
        defRPr=CharacterProperties(sz=850, b=True, solidFill=TEXT,
                                   latin=DrawingFont(typeface=FONT))
    )
    ax.spPr = GraphicalProperties()
    ax.spPr.line = LineProperties(solidFill=AXIS, w=9525)
    ax.majorTickMark = "none"
    ax.minorTickMark = "none"

chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 62000
chart.y_axis.majorUnit = 10000
chart.y_axis.numFmt = "#,##0"
chart.y_axis.majorGridlines = ChartLines()
chart.y_axis.majorGridlines.spPr = GraphicalProperties()
chart.y_axis.majorGridlines.spPr.line = LineProperties(solidFill=GRID, w=9525)
chart.x_axis.majorGridlines = None
chart.x_axis.tickLblPos = "low"

chart.legend.position = "b"
chart.legend.overlay = False
chart.legend.txPr = txpr(size=800)
# a serie tecnica dos rotulos (indice 4) nao aparece na legenda
chart.legend.legendEntry = [LegendEntry(idx=4, delete=True)]

ws.add_chart(chart, "A21")

# impressao: paisagem, ajustada a largura da folha
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.sheet_view.showGridLines = False

OUT = "/home/user/procafe/excel/producao_arabica_robusta.xlsx"

# o Excel recalcula as formulas na abertura (openpyxl nao grava valores em cache)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)


def add_source_linked(path):
    """numFmt de rotulo sem sourceLinked="0" e ignorado (o app usa o formato da
    celula e mostra o numero em vez do texto ON/OFF)."""
    with zipfile.ZipFile(path) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename.startswith("xl/charts/chart"):
                data = re.sub(
                    r'<numFmt formatCode="([^"]*)" />',
                    r'<numFmt formatCode="\1" sourceLinked="0" />',
                    data.decode(),
                ).encode()
            zout.writestr(info, data)


add_source_linked(OUT)
print("ok ->", OUT)
