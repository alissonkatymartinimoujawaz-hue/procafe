# -*- coding: utf-8 -*-
"""Small toolkit on top of python-pptx for a research-style deck.

Every chart is a native (editable) PowerPoint chart. The plot area of each chart
is pinned with a manual layout so that overlay shapes (shaded periods, threshold
lines, call-outs) can be positioned in data coordinates.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION,
                            XL_TICK_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_MARK)
from pptx.oxml.ns import qn
from lxml import etree

# --------------------------------------------------------------------------- palette
C = dict(
    espresso="2B1A12", ink="1F1E1C", ink2="52514E", muted="898781", grid="E1E0D9",
    axis="C3C2B7", panel="F4F3EF", white="FFFFFF", accent="B3261E", crema="D9A441",
    arabica="2A78D6", robusta="EB6834", demand="4A3AA7", supply="1BAF7A",
    deficit="E34948", surplus="A9A79F", grow="4A3AA7", steady="898781", decline="E34948",
    yellow="EDA100", magenta="E87BA4", green="008300", band="EFEDE7",
)
FONT = "Calibri"
SLIDE_W, SLIDE_H = 13.333, 7.5


def rgb(h):
    return RGBColor.from_string(h)


# --------------------------------------------------------------------------- text
def add_text(slide, x, y, w, h, paras, size=14, color=None, bold=False, align="l",
             anchor="t", font=FONT, margin=0.0, italic=False, line_spacing=None,
             fill=None, wrap=True):
    """paras: str | list of paragraphs; a paragraph is str or list of runs;
    a run is str or dict(text, bold, color, size, italic).  A paragraph may
    also be dict(runs=[...], bullet=True, level=0, space_after=4, align='l')."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if fill:
        tb.fill.solid(); tb.fill.fore_color.rgb = rgb(fill)
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        spec = p if isinstance(p, dict) else {"runs": p}
        runs = spec.get("runs", [])
        if isinstance(runs, (str, dict)):
            runs = [runs]
        pal = spec.get("align", align)
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[pal]
        if line_spacing:
            para.line_spacing = line_spacing
        if spec.get("space_after") is not None:
            para.space_after = Pt(spec["space_after"])
        if spec.get("space_before") is not None:
            para.space_before = Pt(spec["space_before"])
        if spec.get("bullet"):
            _bullet(para, spec.get("level", 0), spec.get("bullet_color", C["ink2"]),
                    char=spec.get("char", "•"))
        for r in runs:
            rs = r if isinstance(r, dict) else {"text": r}
            run = para.add_run()
            run.text = rs["text"]
            f = run.font
            f.name = rs.get("font", font)
            f.size = Pt(rs.get("size", spec.get("size", size)))
            f.bold = rs.get("bold", spec.get("bold", bold))
            f.italic = rs.get("italic", italic)
            f.color.rgb = rgb(rs.get("color", spec.get("color", color or C["ink"])))
    return tb


def _bullet(para, level, color, char="•"):
    pPr = para._p.get_or_add_pPr()
    indent = 0.18 + 0.2 * level
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(int(-Inches(0.16))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
        el = pPr.find(qn(tag))
        if el is not None:
            pPr.remove(el)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    s = etree.SubElement(buClr, qn("a:srgbClr")); s.set("val", color)
    buSz = etree.SubElement(pPr, qn("a:buSzPct")); buSz.set("val", "100000")
    buFont = etree.SubElement(pPr, qn("a:buFont")); buFont.set("typeface", "Arial")
    buChar = etree.SubElement(pPr, qn("a:buChar")); buChar.set("char", char)


def add_rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None,
             transparency=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        if transparency is not None:
            _set_alpha(shp.fill, transparency)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(0.75)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    shp.shadow.inherit = False
    return shp


def _set_alpha(fill, transparency_pct):
    srgb = fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int((100 - transparency_pct) * 1000)))


def add_line(slide, x1, y1, x2, y2, color, width=1.0, dash=None):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = getattr(MSO_LINE_DASH_STYLE, dash)
    return ln


# --------------------------------------------------------------------------- deck
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.page = 0

    def save(self, path):
        self.prs.save(path)

    # dark slides -----------------------------------------------------------
    def title_slide(self, kicker, title, subtitle, date_line):
        s = self.prs.slides.add_slide(self.blank)
        self.page += 1
        add_rect(s, 0, 0, SLIDE_W, SLIDE_H, C["espresso"])
        add_text(s, 0.8, 1.55, 11.5, 0.4, kicker.upper(), size=14, color=C["crema"], bold=True)
        add_text(s, 0.8, 2.05, 11.6, 2.2, title, size=40, color=C["white"], bold=True, line_spacing=1.0)
        add_text(s, 0.8, 4.35, 11.2, 1.2, subtitle, size=18, color="E9DFD3")
        add_text(s, 0.8, 6.55, 11.5, 0.4, date_line, size=12, color="BFB2A3")
        return s

    def section_slide(self, number, title, question):
        s = self.prs.slides.add_slide(self.blank)
        self.page += 1
        add_rect(s, 0, 0, SLIDE_W, SLIDE_H, C["espresso"])
        add_text(s, 0.8, 2.2, 3, 1.2, f"{number:02d}", size=66, color=C["crema"], bold=True)
        add_text(s, 0.8, 3.45, 11.5, 1.0, title, size=36, color=C["white"], bold=True)
        add_text(s, 0.8, 4.45, 11.0, 1.2, question, size=18, color="E9DFD3")
        return s

    # content slide ---------------------------------------------------------
    def content_slide(self, section, title, subtitle=None, source=None):
        s = self.prs.slides.add_slide(self.blank)
        self.page += 1
        # section tag (pill) — the deck's recurring motif
        tag_w = 0.16 + 0.083 * len(section)
        pill = add_rect(s, 0.55, 0.32, tag_w, 0.3, C["espresso"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        tf = pill.text_frame
        tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = section.upper()
        r.font.size = Pt(10); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = rgb(C["white"])
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        add_text(s, 0.55, 0.72, 12.2, 0.95, title, size=24, bold=True, color=C["espresso"],
                 anchor="t", line_spacing=0.95)
        if subtitle:
            add_text(s, 0.55, 1.62, 12.2, 0.35, subtitle, size=13, color=C["ink2"])
        # footer
        if source:
            add_text(s, 0.55, 6.98, 11.3, 0.38, "Source: " + source, size=8.5, color=C["muted"],
                     anchor="b", line_spacing=0.95)
        add_text(s, 12.2, 7.05, 0.6, 0.3, str(self.page), size=9, color=C["muted"], align="r", anchor="b")
        return s


# --------------------------------------------------------------------------- charts
def _axis_fmt(axis, color=C["muted"], size=10, num_fmt=None, line=True):
    axis.tick_labels.font.size = Pt(size)
    axis.tick_labels.font.name = FONT
    axis.tick_labels.font.color.rgb = rgb(color)
    if num_fmt:
        axis.tick_labels.number_format = num_fmt
        axis.tick_labels.number_format_is_linked = False
    if line:
        axis.format.line.color.rgb = rgb(C["axis"])
        axis.format.line.width = Pt(0.75)
    else:
        axis.format.line.fill.background()
    axis.major_tick_mark = XL_TICK_MARK.NONE
    axis.minor_tick_mark = XL_TICK_MARK.NONE


def _dpt_no_invert(ser, idx):
    """Some renderers treat a data point without <c:invertIfNegative val=0> as inverted (white fill)."""
    for dpt in ser._element.findall(qn("c:dPt")):
        if dpt.find(qn("c:idx")).get("val") != str(idx):
            continue
        if dpt.find(qn("c:invertIfNegative")) is None:
            el = etree.Element(qn("c:invertIfNegative")); el.set("val", "0")
            dpt.find(qn("c:idx")).addnext(el)


def _no_fill_chart(chart):
    """Transparent chart area + plot area so shapes behind the chart show through."""
    cs = chart._chartSpace
    for parent in (cs, cs.chart.plotArea):
        spPr = parent.find(qn("c:spPr"))
        if spPr is None:
            spPr = etree.SubElement(parent, qn("c:spPr"))
            if parent is cs:
                # c:spPr must come before c:txPr / externalData etc. in chartSpace
                txPr = cs.find(qn("c:txPr"))
                if txPr is not None:
                    txPr.addprevious(spPr)
        for ch in list(spPr):
            spPr.remove(ch)
        etree.SubElement(spPr, qn("a:noFill"))
        ln = etree.SubElement(spPr, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))


def set_plot_layout(chart, x, y, w, h):
    """Pin the inner plot area to fractions of the chart frame."""
    plotArea = chart._chartSpace.chart.plotArea
    layout = plotArea.find(qn("c:layout"))
    if layout is None:
        layout = etree.Element(qn("c:layout"))
        plotArea.insert(0, layout)
    for ch in list(layout):
        layout.remove(ch)
    ml = etree.SubElement(layout, qn("c:manualLayout"))
    for tag, val in (("c:layoutTarget", "inner"), ("c:xMode", "edge"), ("c:yMode", "edge"),
                     ("c:x", x), ("c:y", y), ("c:w", w), ("c:h", h)):
        el = etree.SubElement(ml, qn(tag)); el.set("val", str(val))


class ChartFrame:
    """Keeps the geometry needed to map data coordinates to slide inches."""

    def __init__(self, chart, x, y, w, h, px, py, pw, ph, n_cats=None, vmin=None, vmax=None,
                 xmin=None, xmax=None, horizontal=False):
        self.chart, self.x, self.y, self.w, self.h = chart, x, y, w, h
        self.px, self.py, self.pw, self.ph = px, py, pw, ph
        self.n_cats, self.vmin, self.vmax = n_cats, vmin, vmax
        self.xmin, self.xmax = xmin, xmax
        self.horizontal = horizontal

    # plot rectangle in slide inches
    @property
    def left(self):
        return self.x + self.w * self.px

    @property
    def top(self):
        return self.y + self.h * self.py

    @property
    def width(self):
        return self.w * self.pw

    @property
    def height(self):
        return self.h * self.ph

    def cat_x(self, i):
        return self.left + self.width * (i + 0.5) / self.n_cats

    def cat_edge(self, i):
        return self.left + self.width * i / self.n_cats

    def val_y(self, v):
        return self.top + self.height * (1 - (v - self.vmin) / (self.vmax - self.vmin))

    def xy_x(self, xv):
        return self.left + self.width * (xv - self.xmin) / (self.xmax - self.xmin)


def category_chart(slide, kind, x, y, w, h, categories, series, *, vmin=None, vmax=None,
                   major=None, num_fmt='#,##0', legend=True, legend_pos="t", gap=60, overlap=None,
                   plot=(0.07, 0.1, 0.9, 0.78), cat_font=10, val_font=10, gridlines=True,
                   labels=None, label_fmt=None, label_pos=None, label_size=9, cat_label_skip=None,
                   smooth=False, line_width=2.25, markers=False, val_axis_visible=True,
                   tick_low=False, title=None, title_size=11, reverse_cats=False):
    """kind: 'col', 'col_stacked', 'bar', 'bar_stacked', 'line', 'area_stacked'.
    series: list of dict(name, values, color, [line_dash], [width], [point_colors{i:hex}],
    [labels{i:text}], [marker]).  Returns ChartFrame."""
    ctype = {
        "col": XL_CHART_TYPE.COLUMN_CLUSTERED, "col_stacked": XL_CHART_TYPE.COLUMN_STACKED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED, "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
        "line": XL_CHART_TYPE.LINE_MARKERS if markers else XL_CHART_TYPE.LINE,
        "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    }[kind]
    cd = CategoryChartData()
    cd.categories = categories
    for s in series:
        cd.add_series(s["name"], [None if v is None else float(v) for v in s["values"]])
    gf = slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(10)
    ch.font.color.rgb = rgb(C["ink2"])
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        tp = ch.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size = Pt(title_size); tp.runs[0].font.bold = True
        tp.runs[0].font.color.rgb = rgb(C["ink"]); tp.runs[0].font.name = FONT
    else:
        ch.has_title = False
        atd = ch._chartSpace.chart.find(qn("c:autoTitleDeleted"))
        if atd is not None:
            atd.set("val", "1")
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = {"t": XL_LEGEND_POSITION.TOP, "b": XL_LEGEND_POSITION.BOTTOM,
                              "r": XL_LEGEND_POSITION.RIGHT}[legend_pos]
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10); ch.legend.font.name = FONT
        ch.legend.font.color.rgb = rgb(C["ink2"])
    va, ca = ch.value_axis, ch.category_axis
    _axis_fmt(ca, size=cat_font)
    _axis_fmt(va, size=val_font, num_fmt=num_fmt, line=False)
    if not val_axis_visible:
        va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    va.has_major_gridlines = gridlines
    if gridlines:
        va.major_gridlines.format.line.color.rgb = rgb(C["grid"])
        va.major_gridlines.format.line.width = Pt(0.75)
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    if major:
        va.major_unit = major
    if tick_low:
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if reverse_cats:
        ca.reverse_order = True
    if cat_label_skip:
        cax = ca._element
        for tag in ("c:tickLblSkip",):
            el = cax.find(qn(tag))
            if el is not None:
                cax.remove(el)
        skip = etree.SubElement(cax, qn("c:tickLblSkip")); skip.set("val", str(cat_label_skip))
        # schema order: tickLblSkip must come after c:lblOffset / before c:tickMarkSkip, noMultiLvlLbl
        nml = cax.find(qn("c:noMultiLvlLbl"))
        if nml is not None:
            nml.addprevious(skip)
    plot_obj = ch.plots[0]
    if kind.startswith(("col", "bar")):
        plot_obj.gap_width = gap
        if overlap is not None:
            plot_obj.overlap = overlap
        elif "stacked" in kind:
            plot_obj.overlap = 100
    plot_obj.vary_by_categories = False
    for s, ser in zip(series, plot_obj.series):
        col = s.get("color", C["demand"])
        if kind in ("line",):
            ser.smooth = smooth
            lf = ser.format.line
            lf.color.rgb = rgb(col)
            lf.width = Pt(s.get("width", line_width))
            if s.get("dash"):
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                lf.dash_style = getattr(MSO_LINE_DASH_STYLE, s["dash"])
            if markers or s.get("marker"):
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = s.get("marker_size", 6)
                ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col)
                ser.marker.format.line.color.rgb = rgb(C["white"])
            else:
                ser.marker.style = XL_MARKER_STYLE.NONE
        else:
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
            if kind.startswith(("col", "bar")):
                ser.format.line.color.rgb = rgb(C["white"]); ser.format.line.width = Pt(0.75)
            else:
                ser.format.line.fill.background()
            ser.invert_if_negative = False
        for i, pc in (s.get("point_colors") or {}).items():
            pt = ser.points[i]
            if kind == "line":
                pt.marker.style = XL_MARKER_STYLE.CIRCLE
                pt.marker.size = 7
                pt.marker.format.fill.solid(); pt.marker.format.fill.fore_color.rgb = rgb(pc)
                pt.marker.format.line.color.rgb = rgb(C["white"])
            else:
                pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(pc)
                _dpt_no_invert(ser, i)
        # selective direct labels
        for i, txt in (s.get("labels") or {}).items():
            dl = ser.points[i].data_label
            tf = dl.text_frame
            tf.text = txt
            for para in tf.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(s.get("label_size", label_size))
                    run.font.bold = s.get("label_bold", True)
                    run.font.color.rgb = rgb(s.get("label_color", C["ink"]))
                    run.font.name = FONT
            pos = s.get("label_pos", label_pos)
            if pos:
                dl.position = {"above": XL_LABEL_POSITION.ABOVE, "below": XL_LABEL_POSITION.BELOW,
                               "right": XL_LABEL_POSITION.RIGHT, "left": XL_LABEL_POSITION.LEFT,
                               "out": XL_LABEL_POSITION.OUTSIDE_END, "in": XL_LABEL_POSITION.INSIDE_END,
                               "ctr": XL_LABEL_POSITION.CENTER, "base": XL_LABEL_POSITION.INSIDE_BASE}[pos]
    if labels:  # labels for all points of all series (use sparingly)
        plot_obj.has_data_labels = True
        dls = plot_obj.data_labels
        dls.font.size = Pt(label_size); dls.font.name = FONT; dls.font.color.rgb = rgb(C["ink2"])
        if label_fmt:
            dls.number_format = label_fmt; dls.number_format_is_linked = False
        if label_pos:
            dls.position = {"out": XL_LABEL_POSITION.OUTSIDE_END, "in": XL_LABEL_POSITION.INSIDE_END,
                            "ctr": XL_LABEL_POSITION.CENTER, "above": XL_LABEL_POSITION.ABOVE,
                            "base": XL_LABEL_POSITION.INSIDE_BASE}[label_pos]
    _no_fill_chart(ch)
    set_plot_layout(ch, *plot)
    return ChartFrame(ch, x, y, w, h, *plot, n_cats=len(categories), vmin=vmin, vmax=vmax,
                      horizontal=kind.startswith("bar"))


def scatter_chart(slide, x, y, w, h, series, *, xmin, xmax, ymin, ymax, xmajor=None, ymajor=None,
                  x_fmt='0', y_fmt='0', plot=(0.08, 0.06, 0.88, 0.8), legend=True,
                  marker_size=9, x_title=None, y_title=None, lines=False):
    """series: list of dict(name, points=[(x,y),...], color, labels{i:text}, marker)."""
    cd = XyChartData()
    for s in series:
        ss = cd.add_series(s["name"])
        for (xv, yv) in s["points"]:
            ss.add_data_point(xv, yv)
    ctype = XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS if lines else XL_CHART_TYPE.XY_SCATTER
    gf = slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name = FONT; ch.font.size = Pt(10); ch.font.color.rgb = rgb(C["ink2"])
    ch.has_title = False
    atd = ch._chartSpace.chart.find(qn("c:autoTitleDeleted"))
    if atd is not None:
        atd.set("val", "1")
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10); ch.legend.font.color.rgb = rgb(C["ink2"])
    va, xa = ch.value_axis, ch.category_axis   # for XY, category_axis is the X value axis
    _axis_fmt(va, num_fmt=y_fmt, line=False)
    _axis_fmt(xa, num_fmt=x_fmt, line=True)
    va.minimum_scale, va.maximum_scale = ymin, ymax
    xa.minimum_scale, xa.maximum_scale = xmin, xmax
    if ymajor:
        va.major_unit = ymajor
    if xmajor:
        xa.major_unit = xmajor
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(C["grid"])
    xa.has_major_gridlines = False
    # put the X axis labels at the bottom even when y has negative values
    xa.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    va.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    for s, ser in zip(series, ch.plots[0].series):
        col = s.get("color", C["demand"])
        if lines:
            ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(s.get("width", 2))
            ser.marker.style = XL_MARKER_STYLE.NONE
        else:
            ser.format.line.fill.background()
            ser.marker.style = s.get("marker", XL_MARKER_STYLE.CIRCLE)
            ser.marker.size = s.get("marker_size", marker_size)
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col)
            ser.marker.format.line.color.rgb = rgb(C["white"])
        for i, txt in (s.get("labels") or {}).items():
            dl = ser.points[i].data_label
            dl.text_frame.text = txt
            for para in dl.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(s.get("label_size", 9)); run.font.name = FONT
                    run.font.color.rgb = rgb(C["ink2"]); run.font.bold = False
            dl.position = {"above": XL_LABEL_POSITION.ABOVE, "below": XL_LABEL_POSITION.BELOW,
                           "right": XL_LABEL_POSITION.RIGHT, "left": XL_LABEL_POSITION.LEFT}[
                (s.get("label_pos_map") or {}).get(i, s.get("label_pos", "right"))]
    _no_fill_chart(ch)
    set_plot_layout(ch, *plot)
    return ChartFrame(ch, x, y, w, h, *plot, vmin=ymin, vmax=ymax, xmin=xmin, xmax=xmax)


# --------------------------------------------------------------------------- composites
def stat_block(slide, x, y, w, value, label, value_color=None, size=30, label_size=11):
    add_text(slide, x, y, w, 0.55, value, size=size, bold=True, color=value_color or C["espresso"])
    add_text(slide, x, y + 0.02 + size / 60.0, w, 0.6, label, size=label_size, color=C["ink2"])


def side_panel(slide, x, y, w, h, heading, items, fill=None, heading_color=None, size=12):
    """A light panel with a heading and bullet points (items: list[str|list-of-runs])."""
    add_rect(slide, x, y, w, h, fill or C["panel"])
    paras = [{"runs": [{"text": heading, "bold": True, "color": heading_color or C["espresso"],
                        "size": size + 1}], "space_after": 6}]
    for it in items:
        paras.append({"runs": it if isinstance(it, list) else [it], "bullet": True, "space_after": 5,
                      "size": size})
    add_text(slide, x + 0.18, y + 0.15, w - 0.36, h - 0.3, paras, size=size, color=C["ink"])


def legend_chip(slide, x, y, color, text, size=10, line=False):
    if line:
        add_rect(slide, x, y + 0.09, 0.26, 0.05, color)
    else:
        add_rect(slide, x, y + 0.04, 0.14, 0.14, color)
    add_text(slide, x + (0.32 if line else 0.2), y - 0.02, 3.2, 0.26, text, size=size, color=C["ink2"])


def table(slide, x, y, w, h, header, rows, col_widths=None, font_size=10, header_fill=None,
          cell_fills=None, cell_colors=None, bold_cols=(), align=None, row_height=None):
    """cell_fills / cell_colors: dict {(r,c): hex} for body cells (r from 0)."""
    nr, nc = len(rows) + 1, len(header)
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gt.table
    # remove the default table style banding look
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "1")
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    if row_height:
        for r in tbl.rows:
            r.height = Inches(row_height)

    def _cell(cell, text, bold=False, color=C["ink"], fill=C["white"], al="l", size=font_size):
        cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill)
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = ""
        p = tf.paragraphs[0]
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[al]
        r = p.add_run(); r.text = str(text)
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)

    for c, htxt in enumerate(header):
        _cell(tbl.cell(0, c), htxt, bold=True, color=C["white"], fill=header_fill or C["espresso"],
              al=(align[c] if align else "l"))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            _cell(tbl.cell(r + 1, c), val, bold=(c in bold_cols),
                  color=(cell_colors or {}).get((r, c), C["ink"]),
                  fill=(cell_fills or {}).get((r, c), C["white"] if r % 2 == 0 else "FAFAF8"),
                  al=(align[c] if align else "l"))
    return tbl
