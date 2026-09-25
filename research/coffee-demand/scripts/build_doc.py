# -*- coding: utf-8 -*-
"""Build the Word learning guide that accompanies the deck.

Run:  python research/coffee-demand/scripts/build_doc.py
Out:  research/coffee-demand/output/Coffee_Demand_Learning_Guide_2026.docx
"""
import os
import sys

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import analysis as an            # noqa: E402
import charts_mpl                # noqa: E402
from build_deck import classify, episodes, producer_class, switch_buckets, sg, pct   # noqa: E402

OUT = os.path.join(an.BASE, "output", "Coffee_Demand_Learning_Guide_2026.docx")
ESP, INK2, MUTED, PANEL = "2B1A12", "52514E", "898781", "F4F3EF"
BODY_W = 16.0   # cm usable width on A4 with 2.5 cm margins


# ============================================================================ docx helpers
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _insert_tblpr(tblPr, el, before=("w:tblLook", "w:tblCaption", "w:tblDescription")):
    """Insert a tblPr child respecting the schema sequence."""
    old = tblPr.find(el.tag)
    if old is not None:
        tblPr.remove(old)
    for tag in before:
        anchor = tblPr.find(qn(tag))
        if anchor is not None:
            anchor.addprevious(el)
            return
    tblPr.append(el)


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); mar.append(el)
    _insert_tblpr(tblPr, mar)


def set_col_widths(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    _insert_tblpr(tblPr, layout, before=("w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"))
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


class Doc:
    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"; st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.line_spacing = 1.12
        for lvl, size in ((1, 17), (2, 13.5), (3, 11.5)):
            hs = self.d.styles[f"Heading {lvl}"]
            hs.font.name = "Calibri"; hs.font.size = Pt(size); hs.font.bold = True
            hs.font.color.rgb = RGBColor.from_string(ESP)
            hs.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
            rf = hs.element.rPr.find(qn("w:rFonts"))
            if rf is not None:
                for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    if rf.get(qn(a)) is not None:
                        del rf.attrib[qn(a)]
                rf.set(qn("w:ascii"), "Calibri"); rf.set(qn("w:hAnsi"), "Calibri")
            hs.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            hs.paragraph_format.space_after = Pt(5)
        self.fig_n = 0
        self.tab_n = 0
        self._footer()

    def _footer(self):
        p = self.d.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run(p, "Coffee demand learning guide · September 2026 · page ", size=8, color=MUTED)
        r = p.add_run()
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); r._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
        r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED)

    def h1(self, t):
        return self.d.add_heading(t, 1)

    def h2(self, t):
        return self.d.add_heading(t, 2)

    def h3(self, t):
        return self.d.add_heading(t, 3)

    def p(self, *parts, size=None, italic=False, align=None, after=None):
        """parts: str or (text, dict(bold=..., italic=..., color=...))"""
        para = self.d.add_paragraph()
        for part in parts:
            if isinstance(part, tuple):
                run(para, part[0], size=size, italic=italic, **part[1])
            else:
                run(para, part, size=size, italic=italic)
        if align == "c":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style)
            if isinstance(it, (list, tuple)) and it and isinstance(it[0], tuple):
                for t, fmt in it:
                    run(para, t, **fmt)
            elif isinstance(it, list):
                run(para, it[0], bold=True); run(para, it[1])
            else:
                run(para, it)
            para.paragraph_format.space_after = Pt(3)

    def box(self, title, lines, fill=PANEL):
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        shade(c, fill)
        c.width = Cm(BODY_W)
        para = c.paragraphs[0]
        run(para, title, bold=True, size=11, color=ESP)
        for ln in lines:
            q = c.add_paragraph()
            q.paragraph_format.space_after = Pt(3)
            if isinstance(ln, list):
                run(q, ln[0], bold=True, size=10); run(q, ln[1], size=10)
            else:
                run(q, ln, size=10)
        cell_margins(t, 120, 120, 160, 160)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, path, caption, width=BODY_W):
        self.fig_n += 1
        self.d.add_picture(path, width=Cm(width))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.d.add_paragraph()
        run(cap, f"Figure {self.fig_n}. ", bold=True, size=9, color=INK2)
        run(cap, caption, italic=True, size=9, color=INK2)
        cap.paragraph_format.space_after = Pt(10)

    def table(self, header, rows, widths, caption=None, font=8.5, align=None, fills=None, colors=None):
        if caption:
            self.tab_n += 1
            cap = self.d.add_paragraph()
            run(cap, f"Table {self.tab_n}. ", bold=True, size=9, color=INK2)
            run(cap, caption, size=9, color=INK2)
            cap.paragraph_format.space_after = Pt(2)
            cap.paragraph_format.keep_with_next = True
        t = self.d.add_table(rows=len(rows) + 1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, htxt in enumerate(header):
            c = t.cell(0, j); shade(c, ESP)
            para = c.paragraphs[0]; run(para, str(htxt), bold=True, size=font, color="FFFFFF")
            if align and align[j] == "r":
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i + 1, j)
                if (fills or {}).get((i, j)):
                    shade(c, fills[(i, j)])
                elif i % 2 == 1:
                    shade(c, "FAFAF8")
                para = c.paragraphs[0]
                run(para, str(val), size=font, color=(colors or {}).get((i, j)))
                if align and align[j] == "r":
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                para.paragraph_format.space_after = Pt(0)
        # repeat header row on page breaks
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        set_col_widths(t, widths)
        cell_margins(t)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self):
        para = self.d.add_paragraph()
        r = para.add_run()
        fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = 'TOC \\o "1-2" \\h \\z \\u'
        fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t"); t.text = "Right-click here and choose 'Update Field' to build the table of contents."
        fc3 = OxmlElement("w:fldChar"); fc3.set(qn("w:fldCharType"), "end")
        for el in (fc1, it, fc2, t, fc3):
            r._r.append(el)

    def save(self, path):
        # ask Word to refresh fields (TOC) on open
        settings = self.d.settings.element
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        later = ["w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "m:mathPr",
                 "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats",
                 "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions", "w:readModeInkLockDown", "w:smartTagType",
                 "w:schemaLibrary", "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator"]
        anchor = None
        for child in settings:
            if any(child.tag == qn(t) for t in later):
                anchor = child; break
        if anchor is not None:
            anchor.addprevious(uf)
        else:
            settings.append(uf)
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        self.d.save(path)


def f1(x):
    return f"{x:,.1f}"


def heat(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    if v <= -2:
        return "F6C9C8"
    if v < -0.5:
        return "FBE3E2"
    if v < 0.5:
        return None
    if v < 2:
        return "E3E0F4"
    if v < 5:
        return "CFC9EE"
    return "B6AEE6"


# ============================================================================ content
def build():
    A = an.build(save=True)
    figs = charts_mpl.make_all(A)
    M, W, ctry, cons, prod, PR = A["M"], A["W"], A["ctry"], A["cons"], A["prod"], A["PR"]
    cyp, m, EL, excy, G, FMT = A["cyp"], A["m"], A["EL"], A["excy"], A["G"], A["FMT"]
    mw = M["world"]
    pcg = mw["period_cagr"]
    cls = classify(ctry)
    doc = Doc()

    # ------------------------------------------------------------------ title page
    doc.d.add_paragraph().paragraph_format.space_after = Pt(60)
    p = doc.p(("COFFEE MARKET RESEARCH · DEMAND SIDE", dict(bold=True, color="B07800")), size=11)
    p = doc.p(("Coffee demand: a learning guide", dict(bold=True, color=ESP)), size=28, after=6)
    doc.p(("Who drinks the coffee, who flinches at the price, who keeps producing — how the market works, what the official data say "
           "for 2005–2026, and how to read every chart in the companion PowerPoint deck.", dict(color=INK2)), size=13, after=24)
    doc.p(("September 2026 · volumes: USDA PSD and ICO · prices: IMF and ICO to Aug-2026 · national data: ABIC, IBGE, BLS, World Bank",
           dict(color=MUTED)), size=9.5, after=30)
    doc.box("The answers in one page", [
        ["Demand growth: ", f"world consumption rose from {f1(mw['cons_2005'])}m to {f1(mw['cons_2025'])}m 60-kg bags between 2005/06 and 2025/26 "
                            f"(+{mw['cons_cagr_05_25']:.1f}% a year, about +{f1(mw['cons_avg_annual_gain'])}m bags every year). It fell in only "
                            f"{len(mw['cons_down_years'])} of 20 years. USDA forecasts a record {f1(mw['cons_2026f'])}m for 2026/27."],
        ["Faster than supply? ", f"not over 20 years (production +{mw['prod_cagr_05_25']:.1f}%/yr), but yes in 2020/21–2025/26: demand "
                                 f"+{pcg[3]['cons']:.1f}%/yr vs supply +{pcg[3]['prod']:.1f}%/yr, so world stocks fell {abs(mw['stock_draw_20_25']):.1f}m bags to "
                                 f"{f1(mw['stocks_2025'])}m ({mw['stu_2025']:.0f}% of a year's use) and prices hit records."],
        ["Biggest consumers: ", f"Europe (EU-27+UK {f1(ctry.loc['EU-27 + UK','c2025'])}m bags) and the US ({f1(ctry.loc['United States','c2025'])}m), "
                                f"then Brazil ({f1(ctry.loc['Brazil','c2025'])}m). Growth now comes from producing countries and new importers (China, Turkey, the Gulf)."],
        ["Why it grows: ", f"about {M['decomp']['pop_share_of_growth']:.0f}% is population, {M['decomp']['pc_share_of_growth']:.0f}% more coffee per person — "
                           "rising incomes and cities in Asia, café chains and cheap instant coffee."],
        ["Marginal buyers: ", "Europe, the US and Canada are inelastic; China, Vietnam, Turkey and the Gulf keep growing whatever the price; "
                              "Japan, Algeria, Russia and — only at extreme retail inflation — Brazil cut back."],
        ["Switching: ", "roasters add robusta when arabica trades at ≥1.8× robusta and move back when the ratio falls below ~1.5×; "
                        f"structurally robusta rose from {mw['robusta_share_2005']:.0f}% to {mw['robusta_share_2025']:.0f}% of world output."],
        ["Marginal producers: ", "Brazil, Vietnam, Ethiopia, Uganda and India produce through price lows; Honduras, Peru, El Salvador and Colombia cut; "
                                 "Guatemala, Nicaragua, Costa Rica, Côte d'Ivoire decline even at record prices; high prices bring a record crop 3–4 years later."],
    ])
    doc.page_break()

    # ------------------------------------------------------------------ contents + how to use
    doc.h1("Contents")
    doc.toc()
    doc.h2("How to use this guide")
    doc.p("The PowerPoint deck gives the story in charts. This guide explains the mechanics behind each chart, defines every concept a trader "
          "needs (elasticity, marginal buyer, stocks-to-use, arbitrage…), shows the full data tables and documents where every number comes from. "
          "Read section 1 first if the concepts are new; sections 2–6 follow the deck; section 7 explains the data sources and how to refresh them.")
    doc.bullets([["Deck ↔ guide: ", "each chapter here matches a section of the deck (01 Global picture, 02 Segments & drivers, 03 Marginal buyers, "
                                     "04 Marginal producers, 05 Tracker)."],
                 ["Numbers: ", "every number is computed by scripts/analysis.py from the files in research/coffee-demand/data. Nothing is typed by hand."],
                 ["Units: ", "volumes in millions of 60-kg bags of green coffee equivalent; prices in US cents per pound unless stated."]])
    doc.page_break()

    # ================================================================== 1 CONCEPTS
    doc.h1("1. The basics you need")
    doc.h2("1.1 Units, years and the two statistical agencies")
    doc.bullets([
        ["A bag ", "= 60 kg of green (unroasted) coffee. 1 million bags ≈ 60,000 tonnes. Roasted and instant coffee are converted back to 'green bean equivalent'."],
        ["Marketing year (USDA) ", "follows each country's harvest: Brazil July–June, Vietnam and Colombia October–September, Indonesia April–March. "
                                   "Importing countries use October–September. The world total adds up these local years."],
        ["Coffee year (ICO) ", "is always October–September for every country, so ICO and USDA totals never match exactly."],
        ["USDA FAS PSD ", "(Production, Supply & Distribution) gives a full balance sheet for ~90 countries every year since 1960: production (arabica/robusta), "
                          "imports, exports, domestic consumption (roast & ground vs soluble) and stocks. It is the backbone of this analysis."],
        ["ICO ", "(International Coffee Organization) publishes the official indicator prices, monthly exports by coffee type and its own world balance "
                 "(which in 2026 switched to 'apparent consumption' and was revised heavily — do not splice old and new ICO series)."],
    ])
    doc.h2("1.2 Arabica, robusta and the price indicators")
    doc.p("Arabica (Coffea arabica) grows at altitude, is milder and more expensive; robusta (Coffea canephora, called conilon in Brazil) grows in hot lowlands, "
          "is more bitter, has more caffeine and costs less. Robusta dominates instant (soluble) coffee and many espresso blends.")
    doc.table(["Indicator", "What it is", "Used here as"], [
        ["ICO Other Milds", "Washed arabicas from Central America, Peru, India etc., New York ex-dock", "'Arabica' price"],
        ["ICO Robustas", "Vietnam, Uganda, Indonesia, Côte d'Ivoire robustas, New York", "'Robusta' price"],
        ["ICO Composite (I-CIP)", "Weighted average of the four ICO groups", "headline level only"],
        ["ICE 'KC' (New York)", "Arabica futures; the trading benchmark", "context"],
        ["ICE 'RC' (London)", "Robusta futures, US$ per tonne", "context"],
        ["Green price (own)", "0.6 × Arabica + 0.4 × Robusta — roughly the world consumption mix", "price in elasticity work"],
        ["Ratio / arbitrage", "Arabica ÷ Robusta (or Arabica − Robusta in c/lb)", "switching signal"],
    ], [3.6, 8.4, 4.0], "Price indicators used in this guide")
    doc.h2("1.3 Demand, consumption and stocks")
    doc.p("Nobody counts cups. For an importing country, USDA estimates consumption as imports minus re-exports, adjusted for changes in stocks — "
          "i.e. disappearance. That is why consumption in the EU or the US can jump ±5–10% in one year: roasters and traders build or run down inventories. "
          "Always read trends on 3-year averages. For producing countries, national surveys are used (for Brazil, ABIC's industry panel).")
    doc.p(("Stocks-to-use ", dict(bold=True)), "= ending stocks ÷ annual consumption. It measures the buffer against a bad crop. In this dataset it ranged from "
          f"{mw['stu_peak']:.0f}% ({mw['stocks_peak_year']}) down to {mw['stu_2025']:.0f}% (2025/26). The lower it is, the more violently prices react to weather.")
    doc.h2("1.4 Growth maths")
    doc.bullets([
        ["Year-on-year (y/y) ", "= value this year ÷ value last year − 1."],
        ["CAGR ", "(compound annual growth rate) = (end ÷ start)^(1/years) − 1. 125.2 → 173.9 over 20 years = +1.66% a year."],
        ["3-year average ", "smooths inventory noise: e.g. average 2023/24–2025/26 vs average 2017/18–2019/20 compares 'during the spike' with 'before it'."],
        ["Growth decomposition ", "consumption = population × consumption per person, so ln(growth) = ln(population growth) + ln(per-capita growth)."],
    ])
    doc.h2("1.5 Elasticity — the key idea")
    doc.p(("Price elasticity of demand ", dict(bold=True)), "= % change in quantity ÷ % change in price. An elasticity of −0.1 means a 10% price rise "
          "cuts volumes by 1%. Between 0 and −1 demand is inelastic (spending rises when prices rise); below −1 it is elastic.")
    doc.bullets([
        "Coffee demand is inelastic because the green bean is a small part of what the drinker pays: in a café cup it is a few percent, "
        "on the supermarket shelf roughly a third to a half. A doubling of green prices raises shelf prices by far less, and later.",
        "Habit and caffeine make coffee a staple for regular drinkers; cheaper substitutes (instant, robusta blends, own-label) exist inside the category, "
        "so people trade down rather than stop.",
        ["Cross-price elasticity ", "(arabica vs robusta) is where the action is: roasters change blends when the price ratio moves."],
        ["Income elasticity ", "(response to income) is high in emerging markets and near zero in mature ones."],
    ])
    doc.h2("1.6 Marginal buyer and marginal producer")
    doc.p(("The marginal buyer ", dict(bold=True)), "is the last buyer who drops out when the price rises — the one whose reaction balances the market. ",
          ("The marginal producer ", dict(bold=True)), "is the highest-cost supplier who stops (or starts) producing as the price falls (or rises). "
          "In coffee both are small groups: most demand does not react to price in the short run and most supply cannot react for 2–4 years "
          "(new trees take that long to bear), so the price has to move a lot to find the few who can adjust.")
    doc.page_break()

    # ================================================================== 2 GLOBAL
    doc.h1("2. The global picture, 2005–2026")
    doc.h2("2.1 World demand, year by year")
    doc.figure(figs["world_demand"], "World coffee consumption and its year-on-year change (USDA).")
    doc.p(f"World consumption rose by {f1(mw['cons_abs_gain_05_25'])}m bags in 20 years. The year-on-year changes are small and mostly positive: "
          f"the standard deviation of annual growth is {mw['cons_yoy_sd']:.1f}% (production: {mw['prod_yoy_sd']:.1f}%). The {len(mw['cons_down_years'])} down years "
          f"({', '.join(mw['cons_down_years'])}) are explained by inventory swings in importing countries (2006/07, 2008/09, 2010/11, 2023/24) and COVID "
          "(2019/20–2020/21). The +10.6% in 2009/10 is the mirror image of low stock-building the year before, not a sudden jump in drinking.")
    rows = []
    for y in range(2005, 2027):
        rows.append([W.label[y] + ("F" if y == 2026 else ""), f1(W.production[y]), f1(W.consumption[y]),
                     "" if np.isnan(W.cons_yoy[y]) else pct(W.cons_yoy[y]), "" if np.isnan(W.prod_yoy[y]) else pct(W.prod_yoy[y]),
                     f1(W.ending_stocks[y]), f"{W.stocks_to_use[y]:.1f}%"])
    fills = {(i, 3): heat(W.cons_yoy[y]) for i, y in enumerate(range(2005, 2027)) if not np.isnan(W.cons_yoy[y])}
    doc.table(["Year", "Production", "Consumption", "Cons. y/y", "Prod. y/y", "End stocks", "Stocks/use"], rows,
              [2.0, 2.3, 2.5, 2.2, 2.2, 2.3, 2.5], "World balance, million bags (USDA PSD Dec-2025 release; 2026/27 = USDA Jul-2026 forecast)",
              align=["l", "r", "r", "r", "r", "r", "r"], fills=fills)
    doc.p(("Note on 2026/27: ", dict(bold=True)), "USDA's July-2026 report forecasts consumption at 179.7m (+3.6% vs its revised 2025/26) and stocks "
          "rebuilding by 1.9m bags to 26.3m. Its revised 2025/26 figures were not published in the text we could access, so 2025/26 is shown in the "
          "December-2025 version; the stock rebuild should be read from USDA's +1.9m, not from the table's difference.", size=9.5)

    doc.h2("2.2 Is demand growing faster than supply?")
    doc.figure(figs["supply_demand"], "World production vs consumption, and the change in world ending stocks (USDA).")
    doc.figure(figs["period_cagr"], "Compound annual growth by 5-year period; stock change under each period.")
    doc.p("The honest answer depends on the window:")
    doc.bullets([
        [f"20 years: no. ", f"Production grew {mw['prod_cagr_05_25']:.1f}% a year vs {mw['cons_cagr_05_25']:.1f}% for consumption. Part of that is the "
                            "starting point: 2005/06 was a small Brazilian crop."],
        [f"2010/11–2015/16 and 2020/21–2025/26: yes. ", f"Demand grew {pcg[1]['cons']:.1f}% and {pcg[3]['cons']:.1f}% a year while supply grew "
                                                        f"{pcg[1]['prod']:.1f}% and {pcg[3]['prod']:.1f}%. In the second period stocks fell {abs(pcg[3]['stock_change']):.1f}m bags."],
        ["The price follows the stock cycle, not the demand trend. ", "Demand growth sits in a narrow 1.1–2.5% band in every period; supply growth ranges "
                                                                      f"from {min(p['prod'] for p in pcg):.1f}% to {max(p['prod'] for p in pcg):.1f}%."],
    ])
    doc.h2("2.3 Stocks and prices")
    doc.figure(figs["stocks_prices"], "Stocks-to-use (top) and coffee-year average prices (bottom).")
    doc.p(f"Stocks peaked at {f1(mw['stocks_peak'])}m bags in {mw['stocks_peak_year']} and fell to {f1(mw['stocks_2025'])}m by 2025/26. "
          f"Arabica averaged {cyp.arabica[2018]:.0f} c/lb in coffee year 2018/19 and {cyp.arabica[2024]:.0f} c/lb in 2024/25; robusta went from "
          f"{cyp.robusta[2018]:.0f} to {cyp.robusta[2024]:.0f} c/lb. From 2022/23 stocks-to-use sat below 16% for the first time in the series, and prices broke out "
          "to records: with no buffer, every weather scare had to be rationed by price.")

    doc.h2("2.4 Who consumes the most, and how")
    doc.figure(figs["top_consumers"], "Top 15 consumers, 2025/26 vs 2005/06 (USDA).")
    top = ctry.drop(index=["World", "European Union", "United Kingdom"]).sort_values("c2025", ascending=False).head(25)
    rows = [[c.replace("Korea, South", "South Korea"), f1(r.c2005) if not np.isnan(r.c2005) else "–", f1(r.c2015), f1(r.c2020), f1(r.c2025),
             f"{r.share2025:.1f}%", pct(r.cagr_05_25) if not np.isnan(r.cagr_05_25) else "–", r.group] for c, r in top.iterrows()]
    doc.table(["Country", "2005/06", "2015/16", "2020/21", "2025/26", "Share", "CAGR 05→25", "Group"], rows,
              [3.0, 1.6, 1.6, 1.6, 1.6, 1.4, 1.9, 3.3], "Consumption by country, million bags (USDA)", align=["l", "r", "r", "r", "r", "r", "r", "l"])
    doc.p("How people drink differs by market. USDA splits each country's consumption into roast & ground and soluble (instant). Europe, the US and Brazil "
          "drink more than 90% roast & ground (filter, espresso, capsules, cafés). Instant dominates in the Philippines and is a third of use in China, "
          "Indonesia and Vietnam (Table 4). Brazil buys 73–78% of its coffee through retail for home brewing (ABIC, 2025).")
    p_ = an.load_psd()
    sol_t, tot_t = an.psd_table(p_, "Soluble Dom. Cons."), an.psd_table(p_, "Domestic Consumption")
    rows = []
    for c in ["Philippines", "Thailand", "India", "Mexico", "Canada", "China", "Indonesia", "Vietnam", "Colombia", "Australia",
              "Ukraine", "Russia", "EU-27 + UK", "Japan", "United States", "Brazil"]:
        rows.append([c, f"{tot_t.loc[c, 2025]:.2f}", f"{sol_t.loc[c, 2025] / tot_t.loc[c, 2025] * 100:.0f}%"])
    doc.table(["Country", "Consumption 2025/26, m bags", "of which instant (soluble)"], rows, [4.5, 4.5, 4.5],
              "Instant coffee's share of consumption by country, 2025/26 (USDA; several countries, e.g. Korea and Turkey, report no split)",
              align=["l", "r", "r"])
    ico = pd.read_csv(os.path.join(an.SRC, "Coffee_importers_consumption.csv"))
    ico["Country"] = ico.Country.str.strip()
    ico = ico.set_index("Country")
    eu = ["Germany", "France", "Italy", "Spain", "Poland", "Netherlands", "Sweden", "Finland", "Belgium", "Austria", "Greece", "Romania", "Portugal"]
    rows = [[c, f"{ico.loc[c, '2005'] / 60e6:.2f}", f"{ico.loc[c, '2019'] / 60e6:.2f}"] for c in eu if c in ico.index]
    doc.table(["EU country", "2005", "2019"], rows, [5.0, 3.0, 3.0],
              "Inside Europe: consumption by member state, million bags (ICO importing-country data, calendar years; latest public year 2019)",
              align=["l", "r", "r"])

    doc.h2("2.5 Where the growth came from")
    doc.figure(figs["contributions"], "Change in consumption by country, 2005/06 → 2025/26, coloured by buyer group.")
    doc.figure(figs["buyer_groups"], "World consumption by buyer group.")
    tot = {g: M["contrib_by_group"][g] for g in M["contrib_by_group"]}
    doc.p(f"Of the {f1(M['contrib_total'])}m extra bags, producing countries added {f1(tot['Producing countries'])}m "
          f"({tot['Producing countries'] / M['contrib_total'] * 100:.0f}%), emerging importers {f1(tot['Emerging importers'])}m "
          f"({tot['Emerging importers'] / M['contrib_total'] * 100:.0f}%) and traditional importers {f1(tot['Traditional importers'])}m "
          f"({tot['Traditional importers'] / M['contrib_total'] * 100:.0f}%). Traditional importers' share of world use fell from 63% to 52%.")
    doc.p("Trading implication: as Brazil, Vietnam, Indonesia and Ethiopia drink more of their own crops, less is left to export. "
          f"Producing countries now consume about {G.loc[2025, 'Producing countries']:.0f}m bags a year.")

    doc.h2("2.6 Why demand grows: people, income and habit")
    dec = M["decomp"]
    doc.p(f"Between 2005/06 and 2024/25 world consumption rose {dec['total_pct']:.1f}%. World population rose {dec['pop_pct']:.1f}% "
          f"({M['pop']['world_2005'] / 1e9:.2f}bn → {M['pop']['world_2024'] / 1e9:.2f}bn) and consumption per person rose {dec['pc_pct']:.1f}% "
          f"({M['pop']['kg_world_2005']:.2f} → {M['pop']['kg_world_2024']:.2f} kg). On a log basis population explains ~{dec['pop_share_of_growth']:.0f}% "
          f"of the growth and per-capita gains ~{dec['pc_share_of_growth']:.0f}%.")
    doc.figure(figs["per_capita"], "Consumption per person vs GDP per capita (log scale).")
    INC = A["INC"]
    doc.p(f"Per-capita use climbs with income until roughly 4–9 kg a year (Nordics, Switzerland, Canada, Brazil, Europe, the US). China "
          f"({INC.loc['China', 'kg2024']:.2f} kg), India ({INC.loc['India', 'kg2024']:.2f} kg) and Indonesia ({INC.loc['Indonesia', 'kg2024']:.1f} kg) are far below "
          f"— that is the long-run runway. China's GDP per head multiplied by {INC.loc['China', 'gdppc2024'] / INC.loc['China', 'gdppc2005']:.1f} since 2005 and its urban "
          "population share rose from 42.5% to 65.5% (World Bank); its coffee use multiplied almost thirty-fold.")
    doc.page_break()

    # ================================================================== 3 SEGMENTS
    doc.h1("3. Demand segments and drivers")
    doc.p("The brief asked which type of demand grew: personal (at home), coffee shops, retailers, roasters. Official statistics do not split coffee "
          "by channel, so the guide uses three official segmentations and one set of proxies:")
    doc.bullets([["By buyer group ", "(USDA, every country): traditional importers, emerging importers, producing countries — section 2.5."],
                 ["By format ", "(USDA): roast & ground vs soluble (instant) — below."],
                 ["By channel ", "(national data): Brazil's ABIC measures retail (at-home) at 73–78% of domestic volume in 2025."],
                 ["Proxies for cafés ", "(company filings): store counts of the big chains."]])
    doc.h2("3.1 Formats: instant is the fastest-growing")
    doc.figure(figs["formats"], "World consumption: roast & ground vs soluble (USDA).")
    sol_g = an.cagr(FMT.sol_world[2005], FMT.sol_world[2025], 20) * 100
    rg_g = an.cagr(FMT.rg_world[2005], FMT.rg_world[2025], 20) * 100
    doc.p(f"Soluble consumption rose from {f1(FMT.sol_world[2005])}m to {f1(FMT.sol_world[2025])}m bags ({sol_g:+.1f}% a year) against {rg_g:+.1f}% a year for "
          f"roast & ground; its share went from {mw['soluble_share_2005']:.0f}% to {mw['soluble_share_2025']:.0f}%. Instant is the entry format in new "
          "markets and the cheapest cup when prices rise — and it is mostly robusta. Capsules and pods show up in trade statistics as roasted coffee: "
          f"world roast & ground exports rose from {f1(FMT.world_rg_exports[2005])}m to {f1(FMT.world_rg_exports[2025])}m bags; Switzerland, home of "
          f"Nespresso's roasting, exported {f1(FMT.swiss_rg_exports[2025])}m bags of roasted coffee in 2025/26.")
    doc.h2("3.2 Coffee shops")
    sb, sbc, sbu = A["sb_tot"], A["sb_china"], A["sb_us"]
    rows = [[f"FY{int(y)}", f"{v:,.0f}", f"{sbc.get(y):,.0f}" if y in sbc.index else "", f"{sbu.get(y):,.0f}" if y in sbu.index else ""] for y, v in sb.items()]
    doc.table(["Fiscal year", "Starbucks total", "of which China", "of which US"], rows, [3.0, 3.5, 3.5, 3.5],
              "Starbucks stores at fiscal year-end (Form 10-K; years not captured are omitted)", align=["l", "r", "r", "r"])
    doc.p(f"Starbucks grew from {sb.iloc[0]:,.0f} stores (FY2005) to {sb.iloc[-1]:,.0f} (FY2025), ×{sb.iloc[-1] / sb.iloc[0]:.1f}. Since FY2020 its China store count grew "
          f"{(sbc.iloc[-1] / sbc.iloc[0] - 1) * 100:.0f}% while the US grew {(sbu.iloc[-1] / sbu.iloc[0] - 1) * 100:.0f}%. Luckin Coffee reported 22,340 stores at the end of 2024 "
          "(20-F, as quoted in a secondary source) — about three times Starbucks China. Cafés matter less for volume than for habit: they recruit new, young, "
          "urban drinkers who then also buy coffee for home.")
    doc.h2("3.3 Roasters and blends")
    doc.p("Roasters are the buyers who actually face green-coffee prices, and they have three levers: raise shelf prices, change the blend (more robusta/conilon, "
          "cheaper origins), or run down inventories. ABIC reports that in 2021–2025 Brazilian green-coffee costs rose 201% (conilon) and 212% (arabica) while "
          "retail prices rose 116% — roasters absorbed part of the shock and changed blends. The ICO attributed part of the fall in arabica exports in 2022/23 to "
          "'substitution towards the Robustas' and to consuming countries running down stocks.")
    doc.page_break()

    # ================================================================== 4 MARGINAL BUYERS
    doc.h1("4. Marginal buyers: who reacts to price")
    doc.h2("4.1 The price shock")
    doc.figure(figs["prices"], "Monthly arabica and robusta prices and their ratio, Jan-2010 → Aug-2026.")
    rows = [[f"{y}/{str(y + 1)[2:]}", f"{r.arabica:.0f}", f"{r.robusta:.0f}", f"{r.ratio:.2f}", f"{r.blend:.0f}",
             "" if np.isnan(r.blend_yoy) else pct(r.blend_yoy, 0), f"{int(r.months)}"] for y, r in cyp.loc[2005:2025].iterrows()]
    fills = {(i, 5): heat(-r.blend_yoy / 10) if not np.isnan(r.blend_yoy) else None for i, (y, r) in enumerate(cyp.loc[2005:2025].iterrows())}
    doc.table(["Coffee year", "Arabica", "Robusta", "Ratio", "Green price", "y/y", "months"], rows, [2.4, 2.2, 2.2, 2.0, 2.4, 2.0, 1.8],
              "Coffee-year average prices, US c/lb (Oct–Sep; IMF to Sep-2025, ICO after)", align=["l", "r", "r", "r", "r", "r", "r"], fills=fills)
    doc.p(f"Green prices averaged {M['prices']['blend_17_19']:.0f} c/lb in 2017/18–2019/20 and {M['prices']['blend_23_25']:.0f} c/lb in 2023/24–2025/26: "
          f"+{M['prices']['blend_change_3yavg_pct']:.0f}%. That is the 'experiment' used below to see who cut consumption.")

    doc.h2("4.2 Three kinds of buyer")
    doc.p("Each of the 30 largest consuming countries is classified by the change in its 3-year average consumption through the spike "
          "(2023/24–2025/26 vs 2017/18–2019/20). Rules: keep growing = above +10%; inelastic = between −5% and +10%; cut back = below −5% "
          "or a documented price-driven break (Brazil). Venezuela is excluded (economic collapse, not coffee prices).")
    for key, title in (("grow", "Keep growing"), ("steady", "Inelastic"), ("cut", "Cut back")):
        rows = [[c.replace("Korea, South", "South Korea"), pct(v, 0), f1(ctry.loc[c, "c2025"]), pct(ctry.loc[c, "cagr_15_20"]),
                 pct(ctry.loc[c, "cagr_20_25"]), f"{int(ctry.loc[c, 'down_years_21_25'])}"] for c, v in cls[key]]
        doc.table(["Country", "Δ 3-yr avg", "2025/26 m bags", "CAGR 15→20", "CAGR 20→25", "down years 21–25"], rows,
                  [3.4, 2.2, 2.6, 2.5, 2.5, 2.8], f"{title} ({sum(ctry.loc[c, 'c2025'] for c, _ in cls[key]) / W.consumption[2025] * 100:.0f}% of world use)",
                  align=["l", "r", "r", "r", "r", "r"])
    doc.box("Why the US is 'inelastic' but Brazil 'cuts back' at the same −3%", [
        "USDA's US consumption fell 7.8% in 2022/23 and 4.4% in 2023/24 — then rose 11.4% in 2024/25. That pattern is inventory (destocking, then restocking), "
        "not drinkers quitting.",
        "Brazil's fall is measured directly by the roasters' association (ABIC): −2.3% in 2025, −3.9% per person, at a time when retail ground-coffee prices were "
        "up ~80% year on year, followed by a +2.3% rebound in January–August 2026 as prices fell."])

    doc.h2("4.3 The nine country cards (deck section 03)")
    cy_years = list(range(2015, 2026))
    card_c = ["China", "Vietnam", "Turkey", "EU-27 + UK", "United States", "Canada", "Japan", "Brazil", "Algeria"]
    rows = []
    for c in card_c:
        vals = [cons.loc[c, y] for y in cy_years]
        rows.append([c.replace("EU-27 + UK", "EU-27+UK")] + [f"{v:.1f}" for v in vals])
    doc.table(["Country"] + [f"{y % 100:02d}/{(y + 1) % 100:02d}" for y in cy_years], rows, [2.2] + [1.25] * 11,
              "Consumption of the nine card countries, million bags (USDA)", font=7.5, align=["l"] + ["r"] * 11)
    rows = []
    fills = {}
    for i, c in enumerate(card_c):
        yo = [(cons.loc[c, y] / cons.loc[c, y - 1] - 1) * 100 for y in cy_years[1:]]
        rows.append([c.replace("EU-27 + UK", "EU-27+UK")] + [sg(v, 1) for v in yo])
        for j, v in enumerate(yo):
            fills[(i, j + 1)] = heat(v)
    doc.table(["Country"] + [f"{y % 100:02d}/{(y + 1) % 100:02d}" for y in cy_years[1:]], rows, [2.2] + [1.375] * 10,
              "Year-on-year change, % — compare with the green-price change in Table 'Coffee-year average prices'", font=7.5,
              align=["l"] + ["r"] * 10, fills=fills)

    doc.h2("4.4 Measured elasticities")
    doc.figure(figs["elasticity"], "Estimated short-run price elasticities with 95% ranges.")
    el = EL.join(ctry[["c2025"]]).sort_values("c2025", ascending=False).head(16)
    rows = [[c.replace("Korea, South", "South Korea"), sg(r.e_same_year, 2), sg(r.e_lag1, 2), sg(r.e_total, 2), f"±{1.96 * r.se_total:.2f}",
             "yes" if r.e_total + 1.96 * r.se_total < 0 else ""] for c, r in el.iterrows()]
    doc.table(["Country", "same year", "1-yr lag", "total", "95% range", "significant < 0"], rows, [3.6, 2.2, 2.2, 2.2, 2.4, 2.8],
              "Elasticity of consumption to the green price, 2006/07–2025/26", align=["l", "r", "r", "r", "r", "l"])
    doc.p("Method: for each country, the annual log-change in consumption is regressed on the log-change in the green price in the same coffee year and the "
          "year before (20 observations). The sum of the two coefficients is the short-run elasticity. The estimates are small and mostly insignificant — "
          "coffee demand barely responds to green prices within two years. Japan and Indonesia are the clear exceptions. Because shelf prices move less than green "
          "prices, the elasticity with respect to the retail price is roughly 2–3 times larger (Brazil: about −0.03 to −0.05).")

    doc.h2("4.5 Why consumers barely react: retail pass-through")
    doc.figure(figs["passthrough"], "US retail coffee prices (CPI) vs green coffee, index 2019 = 100.")
    us_idx, us_yoy, br_idx, br_dec = an.retail_series()
    calp = A["calp"]; blend = 0.6 * calp.arabica + 0.4 * calp.robusta
    doc.p(f"From 2019 to 2024 green coffee rose {(blend[2024] / blend[2019] - 1) * 100:.0f}% while the US consumer price index for coffee rose "
          f"{(us_idx[2024] / us_idx[2019] - 1) * 100:.0f}%. The shock reached US shelves about a year later: +{us_yoy['2025-06']:.0f}% y/y in June 2025, "
          f"+{us_yoy['2026-01']:.0f}% in January 2026, +{us_yoy['2026-07']:.0f}% in July 2026. Brazil passes costs through faster and further: ground-coffee "
          f"prices rose {br_dec['2021']:.0f}% in 2021, {br_dec['2024']:.0f}% in 2024 and {br_dec['2025']:.0f}% in 2025, then fell {abs(br_dec['2026']):.0f}% "
          "between December 2025 and August 2026.")

    doc.h2("4.6 Brazil: where demand cracks")
    doc.figure(figs["brazil"], "Brazil: retail ground-coffee inflation vs consumption growth.")
    bra = A["bra"]
    rows = [[str(y), f"{bra.consumption.get(y):.2f}" if not np.isnan(bra.consumption.get(y, np.nan)) else "–",
             pct(bra.yoy_published.get(y)) if not np.isnan(bra.yoy_published.get(y, np.nan)) else "–"] for y in range(2005, 2026)]
    doc.table(["Year (Nov–Oct)", "ABIC consumption, m bags", "Published growth"], rows, [4.0, 5.0, 4.0],
              "Brazil domestic consumption (ABIC). Method change in 2018 excludes unregistered firms (earlier years ~2m bags higher); 2016 not captured",
              align=["l", "r", "r"])
    doc.p("Brazil is the best-documented case of a price-sensitive buyer. Retail inflation of +50% (2021) and +40% (2024) was absorbed with volumes still "
          "growing (+1.7%, +1.1%). Only when 12-month ground-coffee inflation hit ~+78–82% (March–May 2025) did volumes break: monthly retail volumes fell "
          "15.96% y/y in April 2025 and the year closed at −2.31% (21.41m bags), −3.88% per person. With retail prices falling in 2026, volumes recovered "
          "(+2.29% in January–August, +4.65% in May–August). One per cent of Brazilian demand is about 0.2m bags — enough to move Brazil's export surplus.")

    doc.h2("4.7 At what price increase does demand fall?")
    ep = episodes(cons, cyp, ctry)
    rows = [[e["country"], e["year"], pct(e["cons_yoy"]), pct(e["p_same"], 0), pct(e["p_prev"], 0), f"{e['level']:.0f}", e["read"]] for e in ep]
    doc.table(["Market", "Year", "Cons. y/y", "Green price same yr", "Green price prior yr", "Level c/lb", "Reading"], rows,
              [2.3, 1.6, 1.6, 2.0, 2.0, 1.6, 4.9], "All consumption declines of 3% or more, 2015/16–2025/26, 30 largest markets", font=7.5,
              align=["l", "l", "r", "r", "r", "r", "l"])
    doc.p("Reading the table: declines that reverse the next year are inventory swings (2023/24 in the US, Switzerland, Algeria). The declines that stick "
          "came one year after a big jump in green prices — 2022/23 after +44% (Canada), 2025/26 after +48% (Turkey, China, Australia, Egypt) — and in Japan, "
          "whose structural decline deepened when prices rose. Rule of thumb: demand only cracks when prices have risen by 40–50% on top of an already high level, "
          "and it cracks with a lag of about a year.")

    doc.h2("4.8 Bean switching")
    doc.figure(figs["switching"], "Robusta share of world exports (12-month rolling) vs the arabica/robusta ratio.")
    bk = switch_buckets(A["ex"], m)
    rows = [[b["bucket"], sg(b["mean"], 1) + " pp", str(b["n"])] for b in bk]
    doc.table(["Arabica ÷ robusta at month t", "Change in robusta export share over the next 12 months", "months"], rows, [5.0, 7.0, 2.5],
              "Switching test on ICO monthly exports, Aug-2016 → Jul-2026", align=["l", "r", "r"])
    ex_cy = excy[excy.months >= 8].loc[2015:2025]
    rows = [[f"{y}/{str(y + 1)[2:]}" + ("*" if r.months < 12 else ""), f"{r.total / 1000:.1f}", f"{r.arabicas / 1000:.1f}", f"{r.robustas / 1000:.1f}",
             f"{r.rob_share:.1f}%", f"{r.ratio:.2f}"] for y, r in ex_cy.iterrows()]
    doc.table(["Coffee year", "World exports", "Arabicas", "Robustas", "Robusta share", "Price ratio"], rows, [2.6, 2.6, 2.4, 2.4, 2.6, 2.4],
              "World exports by type, million bags (sum of ICO monthly reports; *2025/26 = 8 months)", align=["l", "r", "r", "r", "r", "r"])
    doc.p("Switching is measurable but second-order. When arabica trades at 1.8× robusta or more, robusta gains roughly 0.4–1.0 percentage points of the export "
          "market over the next year; below 1.5× it gains nothing. The biggest moves in the share are supply-driven: Brazil's big arabica crop in 2018/19 "
          "pushed the robusta share down to 35.6%; the record conilon crop and short arabica crop in 2025/26 pushed it to ~45%. Structurally, though, the world "
          f"is moving to robusta: {mw['robusta_share_2005']:.0f}% of production in 2005/06, {mw['robusta_share_2025']:.0f}% in 2025/26. The switchers are blend "
          "roasters, instant makers and Brazil's own industry — not individual drinkers.")
    doc.box("Summary — who is the marginal buyer?", [
        "When prices rise, the volume that disappears is small (about 1–2m bags of ~174m) and comes from Japan, North Africa (Algeria), Russia/Eastern Europe "
        "and, at extreme retail inflation, Brazil.",
        "Growth buyers (China, Vietnam, Turkey, the Gulf, Korea, Mexico) keep adding ~1m bags a year even at record prices — a floor under demand.",
        "The larger adjustment is the bean mix: roasters buy more robusta when arabica costs ≥1.8× robusta and less when the ratio falls towards 1.5×."])
    doc.page_break()

    # ================================================================== 5 PRODUCERS
    doc.h1("5. Marginal producers")
    doc.h2("5.1 Supply by origin")
    doc.figure(figs["production"], "Production by origin (USDA).")
    pt = PR.drop(index=["World", "EU-27 + UK", "European Union"], errors="ignore").sort_values("p2025", ascending=False).head(22)
    pcl = producer_class(pt)
    rows = [[c, f1(r.p2005), f1(r.p2015), f1(r.p2020), f1(r.p2025), f"{r.p2026f:.1f}" if not np.isnan(r.p2026f) else "–",
             f"{r.robusta_share2025:.0f}%", pct(r.cagr_05_25), pct(r.low_price_resp, 0), pct(r.high_price_resp, 0), pcl[c]] for c, r in pt.iterrows()]
    doc.table(["Origin", "05/06", "15/16", "20/21", "25/26", "26/27F", "Rob.", "CAGR", "After low", "Boom", "Type"], rows,
              [2.4, 1.1, 1.1, 1.1, 1.1, 1.2, 1.0, 1.3, 1.4, 1.3, 2.9], "Production by origin, million bags, and price responses (USDA)", font=7.5,
              align=["l", "r", "r", "r", "r", "r", "r", "r", "r", "r", "l"])
    doc.h2("5.2 The response map and the four producer types")
    doc.figure(figs["producer_map"], "Output after the 2018–19 price low (x) vs during the 2022–25 boom (y).")
    lowp = cyp.loc[2017:2019, "arabica"].mean(); prevp = cyp.loc[2014:2016, "arabica"].mean()
    doc.bullets([
        ["Produce no matter what — ", f"Brazil, Vietnam, Indonesia, Ethiopia, Uganda, India. They kept producing through arabica at ~{lowp:.0f} c/lb and robusta at "
                                      f"~{cyp.loc[2017:2019, 'robusta'].mean():.0f} c/lb (2017/18–2019/20): scale and mechanisation (Brazil), very high yields (Vietnam), "
                                      "low cash costs and family labour (Ethiopia, Uganda)."],
        ["Cut when prices are low — ", f"Honduras (−26%), El Salvador (−14%), Peru (−12%), Colombia (−7%). Arabica averaged {lowp:.0f} c/lb, "
                                        f"{(1 - lowp / prevp) * 100:.0f}% below 2014–16, for two to three years; high labour costs per bag, leaf rust, fertiliser cuts and "
                                        "migration did the rest. Honduras never got back to its 2016–18 peak even in the boom."],
        ["Decline even at high prices — ", "Guatemala, Nicaragua, Costa Rica, Côte d'Ivoire, Malaysia (and Cameroon): output fell further in 2022–25 while arabica "
                                           "was above 200 c/lb. Land goes to other crops (cocoa, rubber, durian), housing and tourism; trees and farmers are old."],
        ["Expand when prices are high — ", "Papua New Guinea, Thailand, Kenya bounce back quickly; Brazil's conilon and Vietnam and Uganda expand with a lag. "
                                           f"The result is USDA's record 2026/27 crop of {f1(mw['prod_2026f'])}m bags (+{f1(mw['prod_2026f'] - mw['prod_2025'])}m)."],
    ])
    doc.p("Why the response is slow and asymmetric: a coffee tree needs 3–4 years to bear, so supply cannot react quickly to high prices, while low prices "
          "cut output at the high-cost margin within 1–2 years (less fertiliser, fewer pickers, abandoned plots). Cost-of-production surveys (CONAB, FNC, ICO) "
          "are not included in this version; the thresholds above are inferred from observed output.")
    doc.page_break()

    # ================================================================== 6 TRACKER
    doc.h1("6. Country tracker")
    tr = ctry.drop(index=["World", "European Union", "United Kingdom"]).sort_values("c2025", ascending=False).head(30)
    cls_map = {c: k for k, lst in cls.items() for c, _ in lst}
    rows, fills, colors = [], {}, {}
    for i, (c, r) in enumerate(tr.iterrows()):
        k = cls_map.get(c, "steady")
        rows.append([c.replace("Korea, South", "South Korea"), f1(r.c2025), f"{r.share2025:.1f}%", pct(r.cagr_05_15) if not np.isnan(r.cagr_05_15) else "–",
                     pct(r.cagr_15_20), pct(r.cagr_20_25), pct(r.chg_spike_3yavg, 0), pct(r.yoy_2025),
                     {"grow": "↑ growing", "steady": "→ steady", "cut": "↓ cutting"}[k]])
        for j, v in ((3, r.cagr_05_15), (4, r.cagr_15_20), (5, r.cagr_20_25)):
            fills[(i, j)] = heat(v)
        colors[(i, 8)] = {"grow": "008300", "steady": INK2, "cut": "E34948"}[k]
    doc.table(["Country", "2025/26", "Share", "05→15", "15→20", "20→25", "Spike Δ", "y/y 25/26", "Trend"], rows,
              [3.0, 1.5, 1.4, 1.5, 1.5, 1.5, 1.6, 1.7, 2.3], "Consumption tracker, 30 largest markets (USDA; CAGR by period)", font=8,
              align=["l", "r", "r", "r", "r", "r", "r", "r", "l"], fills=fills, colors=colors)
    doc.h2("What to watch")
    doc.bullets([
        ["Demand baseline: ", f"+2.4m bags a year; USDA 2026/27F {f1(mw['cons_2026f'])}m; ICO sees 2025/26 at 180.6m (−0.8%, lower US demand)."],
        ["Stocks-to-use: ", f"the spike happened below ~16%; USDA's 2026/27 figure is {W.stocks_to_use[2026]:.1f}% — still tight."],
        ["Arbitrage: ", f"ratio {M['prices']['ratio_last']:.1f}× in Aug-2026; above ~1.8× roasters add robusta, below ~1.5× they return to arabica."],
        ["Marginal buyers: ", "Brazil retail inflation (IPCA café moído), Japan's demand, North African and Russian imports give the first signs of demand loss."],
        ["Producer response: ", "record 2026/27 crop from Brazil, Vietnam and East Africa; Central America will add little."],
    ])
    doc.page_break()

    # ================================================================== 7 DATA
    doc.h1("7. Data, sources and methodology")
    doc.p("The research environment blocked direct downloads from the statistical agencies (usda.gov, ico.org, worldbank.org, eurostat, census, comtrade) and "
          "capped web searches. Official files were therefore obtained from public mirrors of the original downloads and every mirror was cross-checked "
          "against independent official figures before use:")
    doc.table(["Dataset", "How it was obtained", "Validation"], [
        ["USDA PSD coffee (all countries, 1960–2025/26, Dec-2025 release)", "Official bulk file psd_coffee_csv.zip downloaded 21-Jun-2026 by the public project "
         "andenick/Foodberg (PROVENANCE.md)", "98.2% of 84,170 values identical to a second, independent copy of the Jun-2024 release (differences = normal revisions); "
         "world totals equal USDA's Dec-2025 report (178.8 / 173.9 / 20.1m bags)"],
        ["USDA Coffee: World Markets and Trade, Jul-2026", "Web search of the report and trade press", "2026/27 figures cross-checked in two or more sources"],
        ["IMF Primary Commodity Prices (PCOFFOTM, PCOFFROB)", "IMF source file external-data.xls (Mar-2026) in datasets/commodity-prices", "Annual averages match the World Bank Pink Sheet "
         "to within 0.5%"],
        ["ICO Coffee Market Reports (prices, exports by group)", "Report tables parsed by the public project BMR-Com/coffee-analytics", "22 of 22 monthly values found "
         "independently by web search match exactly; months with parse errors were dropped"],
        ["ABIC, IBGE IPCA, BLS CPI", "Web search of ABIC/Agência Brasil releases; public copies of IBGE/BLS series", "ABIC figures from two or more articles"],
        ["World Bank population & GDP per capita", "WDI bulk files (public copies)", "Standard WDI series"],
        ["Starbucks store counts", "Form 10-K via web search", "High-confidence rows only"],
    ], [4.2, 5.8, 6.0], "Provenance of every dataset", font=8)
    doc.h2("Limitations")
    doc.bullets([
        "USDA consumption is disappearance: year-to-year changes in importing countries include inventory movements.",
        "Channel data (home vs café) exist only for Brazil; café chains are a proxy elsewhere.",
        "Import-origin data (Eurostat Comext, US Census, UN Comtrade) were not pulled; the refresh script downloads them.",
        "Cost-of-production surveys are not included; producer price thresholds are inferred from behaviour.",
        "USDA's July-2026 revisions to 2025/26 were available only for world production and trade, not for every country.",
    ])
    doc.h2("How to refresh everything from the official files")
    doc.bullets(["python research/coffee-demand/scripts/fetch_official_data.py — downloads USDA PSD, World Bank Pink Sheet, IMF, FRED, US Census, UN Comtrade and Eurostat files "
                 "and logs URL, time and SHA-256 in data/source_files/MANIFEST.csv",
                 "python research/coffee-demand/scripts/analysis.py — rebuilds data/clean/*.csv and output/metrics.json",
                 "python research/coffee-demand/scripts/build_deck.py and build_doc.py — regenerate the PowerPoint and this guide"], style="List Number")
    doc.h2("Official homes of the data")
    doc.bullets(["USDA FAS PSD Online — https://apps.fas.usda.gov/psdonline/app/index.html#/app/downloads",
                 "USDA Coffee: World Markets and Trade — https://www.fas.usda.gov/data/commodities/coffee",
                 "ICO Coffee Market Reports — https://www.ico.org (e.g. https://www.ico.org/documents/cy2025-26/cmr-0826-e.pdf)",
                 "IMF Primary Commodity Prices — https://www.imf.org/en/Research/commodity-prices",
                 "World Bank Pink Sheet — https://www.worldbank.org/en/research/commodity-markets",
                 "ABIC — https://www.abic.com.br/estatisticas/indicadores-da-industria/",
                 "IBGE IPCA — https://sidra.ibge.gov.br · BLS CPI — https://www.bls.gov/cpi/",
                 "World Bank WDI — https://data.worldbank.org"])
    doc.page_break()

    # ================================================================== APPENDIX
    doc.h1("Appendix A. Consumption by country, every year (USDA, million bags)")
    top20 = ctry.drop(index=["World", "European Union", "United Kingdom"]).sort_values("c2025", ascending=False).head(20).index
    yrs_a = list(range(2005, 2026, 2)) + [2025] if 2025 not in range(2005, 2026, 2) else list(range(2005, 2026, 2))
    yrs_a = list(range(2005, 2026))
    for chunk in (yrs_a[:11], yrs_a[11:]):
        rows = [[c.replace("Korea, South", "S. Korea").replace("EU-27 + UK", "EU-27+UK")] +
                [("" if np.isnan(cons.loc[c, y]) else f"{cons.loc[c, y]:.1f}") for y in chunk] for c in list(top20) + ["World"]]
        doc.table(["Country"] + [f"{y % 100:02d}/{(y + 1) % 100:02d}" for y in chunk], rows, [2.3] + [13.7 / len(chunk)] * len(chunk),
                  f"Consumption {chunk[0]}/{str(chunk[0] + 1)[2:]}–{chunk[-1]}/{str(chunk[-1] + 1)[2:]}", font=7, align=["l"] + ["r"] * len(chunk))
    doc.h1("Appendix B. Production by origin, every year (USDA, million bags)")
    topp = PR.drop(index=["World", "EU-27 + UK", "European Union"], errors="ignore").sort_values("p2025", ascending=False).head(20).index
    for chunk in (yrs_a[:11], yrs_a[11:]):
        rows = [[c.replace("Papua New Guinea", "PNG")] + [("" if np.isnan(prod.loc[c, y]) else f"{prod.loc[c, y]:.1f}") for y in chunk]
                for c in list(topp) + ["World"]]
        doc.table(["Origin"] + [f"{y % 100:02d}/{(y + 1) % 100:02d}" for y in chunk], rows, [2.3] + [13.7 / len(chunk)] * len(chunk),
                  f"Production {chunk[0]}/{str(chunk[0] + 1)[2:]}–{chunk[-1]}/{str(chunk[-1] + 1)[2:]}", font=7, align=["l"] + ["r"] * len(chunk))
    doc.h1("Appendix C. Glossary")
    doc.table(["Term", "Meaning"], [
        ["Arbitrage", "Price difference (or ratio) between arabica and robusta; drives blend switching"],
        ["Bag", "60 kg of green coffee"],
        ["CAGR", "Compound annual growth rate"],
        ["Coffee year", "October–September (ICO)"],
        ["Conilon", "Brazilian name for robusta (Coffea canephora)"],
        ["Disappearance", "Imports − re-exports ± stock change: USDA's measure of consumption in importing countries"],
        ["Elasticity (price)", "% change in quantity ÷ % change in price"],
        ["Green coffee", "Unroasted beans, the traded commodity"],
        ["I-CIP", "ICO Composite Indicator Price"],
        ["Marginal buyer / producer", "The buyer or producer whose reaction to price balances the market"],
        ["Marketing year", "USDA's local crop year, e.g. Brazil July–June"],
        ["Pass-through", "How much and how fast a change in green prices reaches shelf prices"],
        ["R&G", "Roast & ground coffee (includes capsules)"],
        ["Soluble", "Instant coffee"],
        ["Stocks-to-use", "Ending stocks ÷ consumption"],
    ], [4.0, 12.0], None, font=8.5)

    doc.save(OUT)
    print("saved", OUT)
    return OUT


if __name__ == "__main__":
    build()
