#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Extract the annual coffee series (production, area, yield, trees) of every
origin in the "Balance Sheet" workbook into one tidy CSV.

    python timeseries/extract_series.py "path/to/Balance Sheet.xlsx"

Output -> timeseries/data/coffee_series.csv with columns
    country, source, variable, series, unit, crop_year, year, value
and timeseries/data/exog_brazil.csv: Brazil ON/OFF flag of each crop year and
annual rainfall (mm) / mean temperature (degC) by state (Minas Gerais,
Espirito Santo, Sao Paulo, Parana, others) plus production-weighted averages
for arabica, robusta and total, used by armax_models.py.

* `year` is the first calendar year of the crop year ("2024/2025" -> 2024).
* Only "main" crop-year columns are read.  Revision columns of the same crop
  year ("2024/2025 May", "23/24 (new post)", ...) are skipped so that every
  series has exactly one value per year.
* Excel error strings (#DIV/0!, #REF!), zeros and negative values are treated as missing:
  a national coffee area, tree stock or production is never really zero,
  the sheets use 0 as an empty placeholder.
* Sheets that are an exact copy of another sheet (Ethiopia and Peru are, at
  the time of writing, copies of the Honduras block) are reported and skipped:
  modelling them would only repeat the Honduras results under another name.
"""
import argparse
import csv
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "data", "coffee_series.csv")

YEAR_RE = re.compile(r"^\s*(\d{4})\s*/\s*(\d{4})\s*(\((?:ON|OFF)\))?\s*$", re.I)

# A "row spec" is (variable, series label, unit, row number) or, for a series
# computed from two rows, (variable, label, unit, ("sum", row_a, row_b)) or
# ("ratio", row_num, row_den, factor).
COUNTRIES = {
    "Brazil": {
        "sheet": "Brazil", "source": "USDA / CONAB", "crop_year_start": "July",
        "rows": [
            # rows 54-56 (million bags, USDA marketing years) rather than rows 27-28:
            # rows 27-28 are shifted one crop year late for 2001/02-2008/09
            # (their "2003/2004 (OFF)" holds the 2002/03 record crop).
            ("production", "Production Arabica", "1000 60-kg bags", ("scale", 55, 1000.0)),
            ("production", "Production Robusta", "1000 60-kg bags", ("scale", 56, 1000.0)),
            ("production", "Production Total",   "1000 60-kg bags", ("scale", 54, 1000.0)),
            ("area", "Area total",        "1000 ha", 4),
            ("area", "Area non-bearing",  "1000 ha", 5),
            ("area", "Area bearing",      "1000 ha", 6),
            ("trees", "Trees non-bearing", "million trees", 8),
            ("trees", "Trees bearing",     "million trees", 14),
            ("trees", "Trees total",       "million trees", 15),
            # bags per bearing hectare from the aligned production block
            ("yield", "Yield (production / bearing area)", "bags/ha", ("ratio", ("scale", 54, 1000.0), 6, 1.0)),
        ],
    },
    "Colombia": {
        "sheet": "Colombia W.A", "source": "USDA", "crop_year_start": "October",
        "rows": [
            ("production", "Production Arabica", "1000 60-kg bags", 21),
            ("area", "Area total",       "1000 ha", 4),
            ("area", "Area non-bearing", "1000 ha", 5),
            ("area", "Area bearing",     "1000 ha", 6),
            ("trees", "Trees total",       "million trees", 7),
            ("trees", "Trees non-bearing", "million trees", 10),
            ("trees", "Trees bearing",     "million trees", 15),
            ("yield", "Yield", "bags/ha", 24),
        ],
    },
    "Colombia (Fedecafe)": {
        "sheet": "Colombia W.A", "source": "Fedecafe / Agronet", "crop_year_start": "January",
        "rows": [
            ("production", "Production Arabica", "1000 60-kg bags", 48),
            ("area", "Area total",       "1000 ha", 27),
            ("area", "Area non-bearing", "1000 ha", 28),
            ("area", "Area bearing",     "1000 ha", 29),
            ("trees", "Trees total", "million trees", 31),
            ("yield", "Yield", "bags/ha", 49),
        ],
    },
    "Honduras": {
        "sheet": "Honduras", "source": "USDA", "crop_year_start": "October",
        "rows": [
            ("production", "Production Arabica", "1000 60-kg bags", 24),
            ("area", "Area total",       "1000 ha", 3),
            ("area", "Area non-bearing", "1000 ha", 4),
            ("area", "Area bearing",     "1000 ha", 5),
            ("trees", "Trees non-bearing", "million trees", 8),
            ("trees", "Trees bearing",     "million trees", 13),
            ("trees", "Trees total",       "million trees", ("sum", 8, 13)),
            ("yield", "Yield", "bags/ha", 25),
        ],
    },
    "Ethiopia": {"sheet": "Ethiopia", "same_as": "Honduras"},
    "Peru":     {"sheet": "Peru",     "same_as": "Honduras"},
    "Indonesia": {
        "sheet": "Indonesia", "source": "USDA", "crop_year_start": "April",
        "rows": [
            ("production", "Production Robusta", "1000 60-kg bags", 24),
            ("production", "Production Arabica", "1000 60-kg bags", 25),
            ("production", "Production Total",   "1000 60-kg bags", ("sum", 24, 25)),
            ("area", "Area non-bearing", "million ha", 4),
            ("area", "Area bearing",     "million ha", 5),
            ("area", "Area total",       "million ha", ("sum", 4, 5)),
            ("trees", "Trees non-bearing", "million trees", 8),
            ("trees", "Trees bearing",     "million trees", 13),
            ("trees", "Trees total",       "million trees", ("sum", 8, 13)),
            # bags per bearing hectare = production (1000 bags) / (bearing area (million ha) * 1000)
            ("yield", "Yield (production / bearing area)", "bags/ha", ("ratio", ("sum", 24, 25), 5, 1.0 / 1000.0)),
            ("production", "Production (BPS)", "1000 t", 33),
            ("area", "Area total (BPS)", "million ha", 32),
        ],
    },
    "Vietnam": {
        "sheet": "Vietnam", "source": "USDA", "crop_year_start": "November",
        "rows": [
            ("production", "Production Robusta", "1000 60-kg bags", 24),
            ("production", "Production Arabica", "1000 60-kg bags", 25),
            ("production", "Production Total",   "1000 60-kg bags", ("sum", 24, 25)),
            ("area", "Area total",       "1000 ha", 3),
            ("area", "Area non-bearing", "1000 ha", 4),
            ("area", "Area bearing",     "1000 ha", 5),
            ("yield", "Yield", "t/ha", 26),
        ],
    },
    "Uganda": {
        "sheet": "Uganda", "source": "USDA", "crop_year_start": "October",
        "rows": [
            ("production", "Production Arabica", "1000 60-kg bags", 12),
            ("production", "Production Robusta", "1000 60-kg bags", 13),
            ("production", "Production Total",   "1000 60-kg bags", 14),
            ("area", "Area total",       "1000 ha", 3),
            ("area", "Area non-bearing", "1000 ha", 4),
            ("area", "Area bearing",     "1000 ha", 5),
            ("yield", "Yield (per harvested ha)", "bags/ha", 15),
            ("yield", "Output per planted ha", "bags/ha", 18),
        ],
    },
}


# Brazil: ON/OFF label of the crop year (row 1) and state weather rows
# (annual rainfall in mm and mean temperature in degC, rows under each state).
BRAZIL_WEATHER = {          # state: (rain row, temp row)
    "mg": (36, 37), "es": (41, 42), "sp": (44, 45), "pr": (47, 48), "others": (52, 53),
}
BRAZIL_SHARES = {           # production rows (million bags) used as weather weights
    "arabica": {"mg": 32, "es": 39, "sp": 43, "pr": 46, "others": 50},
    "robusta": {"es": 40, "others": 51},
}
EXOG_OUT = os.path.join(HERE, "data", "exog_brazil.csv")


def extract_brazil_exog(wb, path):
    ws = wb["Brazil"]
    cols = year_columns(ws)
    labels = {}
    for c in ws[1]:
        if isinstance(c.value, str):
            m = YEAR_RE.match(c.value)
            if m and m.group(3):
                labels[int(m.group(1))] = 1 if "ON" in m.group(3).upper() else 0
    years = [y for _, _, y in cols]
    # ON/OFF before the first label: the cycle alternates (2002/03 was the ON record crop)
    first = min(labels)
    for y in years:
        if y not in labels:
            labels[y] = labels[first] if (first - y) % 2 == 0 else 1 - labels[first]
    weather = {st: (row_values(ws, r, cols), row_values(ws, t, cols)) for st, (r, t) in BRAZIL_WEATHER.items()}
    shares = {}
    for grp, rows in BRAZIL_SHARES.items():
        vals = {st: row_values(ws, r, cols) for st, r in rows.items()}
        mean = {st: sum(v for v in d.values() if v) / max(1, sum(1 for v in d.values() if v)) for st, d in vals.items()}
        tot = sum(mean.values())
        shares[grp] = {st: m / tot for st, m in mean.items()}
    shares["total"] = {}
    for grp, w in (("arabica", 0.68), ("robusta", 0.32)):   # long-run arabica / robusta split
        for st, sh in shares[grp].items():
            shares["total"][st] = shares["total"].get(st, 0) + w * sh
    fields = ["year", "crop_year", "onoff"] + [f"{k}_{st}" for st in BRAZIL_WEATHER for k in ("rain", "temp")] \
             + [f"{k}_{g}" for g in ("arabica", "robusta", "total") for k in ("rain", "temp")]
    rows_out = []
    for ci, crop_year, y in cols:
        rec = {"year": y, "crop_year": crop_year, "onoff": labels[y]}
        for st, (rain, temp) in weather.items():
            rec[f"rain_{st}"], rec[f"temp_{st}"] = rain[y], temp[y]
        for grp, sh in shares.items():
            for k in ("rain", "temp"):
                vals = [(rec[f"{k}_{st}"], w) for st, w in sh.items() if rec[f"{k}_{st}"] is not None]
                rec[f"{k}_{grp}"] = round(sum(v * w for v, w in vals) / sum(w for _, w in vals), 2) if vals else None
        if any(rec[f] is not None for f in fields[3:]):
            rows_out.append(rec)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows_out)
    print(f"Brazil ON/OFF + weather: {len(rows_out)} crop years -> {path}")
    print("  weather weights:", {g: {st: round(v, 2) for st, v in sh.items()} for g, sh in shares.items()})


def year_columns(ws):
    """Return [(col_index, crop_year_label, first_year)] for main crop-year headers."""
    cols = []
    for c in ws[1]:
        if not isinstance(c.value, str):
            continue
        m = YEAR_RE.match(c.value)
        if not m:
            continue
        y1, y2 = int(m.group(1)), int(m.group(2))
        if y2 != y1 + 1:          # e.g. the "2023/2025 May" typo column
            continue
        cols.append((c.column, f"{y1}/{y2}", y1))
    return cols


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if v <= 0 else float(v)     # 0 / negative = empty placeholder or broken formula
    return None


def row_values(ws, row, cols):
    return {y: num(ws.cell(row=row, column=ci).value) for ci, _, y in cols}


def spec_values(ws, spec, cols):
    if isinstance(spec, int):
        return row_values(ws, spec, cols)
    kind = spec[0]
    if kind == "sum":
        a, b = spec_values(ws, spec[1], cols), spec_values(ws, spec[2], cols)
        return {y: (a[y] + b[y]) if (a[y] is not None and b[y] is not None) else None for y in a}
    if kind == "ratio":
        n, d, k = spec_values(ws, spec[1], cols), spec_values(ws, spec[2], cols), spec[3]
        return {y: (n[y] / d[y] * k) if (n[y] is not None and d[y] is not None) else None for y in n}
    if kind == "scale":
        a, k = spec_values(ws, spec[1], cols), spec[2]
        return {y: (a[y] * k) if a[y] is not None else None for y in a}
    raise ValueError(spec)


def sheet_fingerprint(ws, rows, max_cols=60):
    """Numeric cells of the modelled rows (labels in column A excluded)."""
    return tuple(tuple(num(ws.cell(row=r, column=c).value) for c in range(2, max_cols + 1))
                 for r in sorted(rows))


def spec_rows(spec):
    if isinstance(spec, int):
        return {spec}
    if spec[0] == "sum":
        return spec_rows(spec[1]) | spec_rows(spec[2])
    if spec[0] == "ratio":
        return spec_rows(spec[1]) | spec_rows(spec[2])
    if spec[0] == "scale":
        return spec_rows(spec[1])
    raise ValueError(spec)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="Balance Sheet workbook")
    ap.add_argument("--out", default=OUT)
    ap.add_argument("--exog-out", default=EXOG_OUT)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    fingerprints = {}
    records = []
    skipped = []

    for country, cfg in COUNTRIES.items():
        ws = wb[cfg["sheet"]]
        if "same_as" in cfg:
            # Same row layout as the reference sheet; skip it while its numbers
            # are still a copy of the reference sheet.
            ref = COUNTRIES[cfg["same_as"]]
            cfg = dict(ref, sheet=cfg["sheet"])
            if sheet_fingerprint(ws, cfg["_rows"]) == fingerprints.get(ref["sheet"]):
                skipped.append((country, ref["sheet"]))
                continue
        cfg["_rows"] = set().union(*(spec_rows(spec) for _, _, _, spec in cfg["rows"]))
        fingerprints[cfg["sheet"]] = sheet_fingerprint(ws, cfg["_rows"])

        cols = year_columns(ws)
        for variable, label, unit, spec in cfg["rows"]:
            vals = spec_values(ws, spec, cols)
            n = 0
            for ci, crop_year, y in cols:
                v = vals[y]
                if v is None:
                    continue
                n += 1
                records.append(dict(country=country, source=cfg["source"], variable=variable,
                                    series=label, unit=unit, crop_year=crop_year, year=y,
                                    value=round(v, 4)))
            print(f"{country:22s} {label:36s} {n:3d} obs")

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with open(args.out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["country", "source", "variable", "series", "unit",
                                          "crop_year", "year", "value"])
        w.writeheader()
        w.writerows(records)
    print(f"\n{len(records)} values written to {args.out}")
    extract_brazil_exog(wb, args.exog_out)
    for country, ref in skipped:
        print(f"WARNING: sheet '{country}' is an exact copy of '{ref}' -> skipped. "
              f"Fill it with {country} data and re-run to model it.", file=sys.stderr)


if __name__ == "__main__":
    main()
