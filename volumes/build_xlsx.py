#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build volumes/suivi_cecafe.xlsx from volumes/cecafe_daily.csv (real data only).

Sheets
  Definitions : how to read Cecafé's Resumo Diário (certificate / customs / shipment ...)
  Donnees     : raw rows (one per day × table × unit), the same layout as cecafe_daily.csv
  Arabica, Robusta : interactive dashboards driven by the "Date affichée" cell:
      1. per stage, all ports: day D (red), month-to-date (green), % change vs previous month (blue, right axis)
      2. per port × stage: movement of the day
      3. per port × stage: % change of the month-to-date vs previous month (bars up / down)
      4. per port × stage: month-to-date vs previous month same period (supporting numbers)
      5. history: every day collected, per stage, with % change vs previous month

Everything is formulas over the Donnees sheet: paste new rows there and the dashboards follow.
Run:  python volumes/build_xlsx.py      (needs openpyxl)
"""
import csv, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties
from openpyxl.formatting.rule import CellIsRule

MONTHS_FR = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
def month_label(date_cell, offset=0):
    """Excel formula giving e.g. 'Septembre 2026' for the month of date_cell (offset=-1 -> previous month)."""
    d = date_cell if offset == 0 else f"DATE(YEAR({date_cell}),MONTH({date_cell})+({offset}),1)"
    names = ",".join(f'"{m}"' for m in MONTHS_FR)
    return f'=CHOOSE(MONTH({d}),{names})&" "&YEAR({d})'


HERE = os.path.dirname(os.path.abspath(__file__))
CSV = os.path.join(HERE, "cecafe_daily.csv")
OUT = os.path.join(HERE, "suivi_cecafe.xlsx")
NMAX = 5000                       # rows reserved in Donnees for future days
F = "Arial"

def font(**k): k.setdefault("name", F); k.setdefault("size", 10); return Font(**k)
H = PatternFill("solid", fgColor="6F4E37"); HF = font(bold=True, color="FFFFFF")
INP = PatternFill("solid", fgColor="FFFF00"); BLUE = font(color="0000FF"); SUB = PatternFill("solid", fgColor="F3EBE2")
thin = Side(style="thin", color="D9D9D9"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
RED, RED2, RED3 = "C0392B", "E07B6B", "F2B8B0"
BLU, BLU2, BLU3 = "1F4E9C", "2F6FBF", "9DC3F0"
GRN = "2E8B57"

STAGES = [("certificates", "Certificats d'origine"), ("customs", "Dédouanement"), ("shipment", "Embarquement")]
PORTS = [("santos", "Santos"), ("vitoria", "Vitória (ES)"), ("rio", "Rio de Janeiro"),
         ("minas", "REDEX / EADI Minas Gerais"), ("salvador", "Salvador"), ("others", "Autres")]
TYPES = {"Arabica": ("E", "I", "M"), "Robusta": ("F", "J", "N")}   # day, cum, prev columns in Donnees

def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.fill = H; c.font = HF; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[L(i)].width = x

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = font(bold=True, size=14, color="6F4E37")
    if sub: ws["A2"] = sub; ws["A2"].font = font(italic=True, color="7A6F66")

def color_bar(s, rgb): s.graphicalProperties.solidFill = rgb; s.graphicalProperties.line.solidFill = rgb
def color_line(s, rgb, width=22000):
    s.graphicalProperties.line.solidFill = rgb; s.graphicalProperties.line.width = width
    s.marker.symbol = "circle"; s.marker.size = 6
    s.marker.graphicalProperties.solidFill = rgb; s.marker.graphicalProperties.line.solidFill = rgb
    s.smooth = False


def read_rows():
    if not os.path.exists(CSV):
        return []
    rows = list(csv.DictReader(open(CSV, encoding="utf-8", newline="")))
    rows.sort(key=lambda r: (r["date"], r["table"], r["unit"]))
    return rows


def sheet_definitions(wb):
    ws = wb.active; ws.title = "Definitions"
    title(ws, "Comment lire le « Resumo Diário » de Cecafé", "Cecafé = conseil des exportateurs de café du Brésil. Unité : sacs de 60 kg. Une publication par jour ouvré.")
    header(ws, 4, ["Étape (tableau de la page)", "Ce que c'est", "Comment le lire"])
    defs = [
        ("1. Certificats d'origine émis", "Le certificat est émis quand un exportateur a une VENTE conclue et un lot prêt, en entrepôt. Le café est vendu, la paperasse export existe, mais il n'a pas bougé.",
         "Signal le plus précoce. « Vente faite, café prêt. » Indique la demande et ce qui va partir dans les 1 à 3 semaines."),
        ("2. Dédouanement (customs clearance)", "Le lot est passé en douane au bureau indiqué (port ou entrepôt douanier intérieur comme REDEX/EADI Minas). Il est libéré pour partir.",
         "« Prêt à partir. » Le total dédouané du jour est égal au total embarqué : c'est le même café vu par le bureau de douane."),
        ("3. Embarquement (shipment)", "Le lot est chargé sur le navire ou le camion. Sortie physique réelle du pays.",
         "« Parti pour de vrai. » Vu par le point d'embarquement : Santos, Rio, autres. Le café dédouané à Vitória ou dans le Minas part le plus souvent par Santos."),
    ]
    r = 5
    for a, b, c in defs:
        ws.cell(row=r, column=1, value=a).font = font(bold=True)
        ws.cell(row=r, column=2, value=b); ws.cell(row=r, column=3, value=c)
        for j in range(1, 4):
            ws.cell(row=r, column=j).alignment = Alignment(wrap_text=True, vertical="top"); ws.cell(row=r, column=j).border = BOX
        ws.row_dimensions[r].height = 62; r += 1
    r += 1
    header(ws, r, ["Bloc de colonnes", "Signification", "Comment le lire"]); r += 1
    cols = [
        ("Mouvement du jour", "Sacs traités ce jour ouvré.", "À comparer à la veille et à la moyenne des jours précédents. Un lundi ou une fin de mois sont plus chargés."),
        ("Cumul (acumulado)", "Total depuis le 1er du mois jusqu'au jour affiché.", "En fin de mois, c'est le total mensuel, à vérifier avec le rapport mensuel Cecafé."),
        ("Mois précédent (mês anterior)", "Même période du mois précédent : les mêmes premiers jours, pas le mois entier.", "Variation % = cumul ÷ mois précédent − 1. Positif = rythme plus rapide que le mois dernier."),
    ]
    for a, b, c in cols:
        ws.cell(row=r, column=1, value=a).font = font(bold=True)
        ws.cell(row=r, column=2, value=b); ws.cell(row=r, column=3, value=c)
        for j in range(1, 4):
            ws.cell(row=r, column=j).alignment = Alignment(wrap_text=True, vertical="top"); ws.cell(row=r, column=j).border = BOX
        ws.row_dimensions[r].height = 46; r += 1
    r += 1
    header(ws, r, ["Lecture", "Ce que ça dit", ""]); r += 1
    reads = [
        ("Certifiés − embarqués (cumul)", "Café vendu à l'export mais pas encore parti : il attend dans les entrepôts portuaires et intérieurs. S'il monte, la file d'attente grossit ; s'il baisse, les embarquements rattrapent."),
        ("Certificats en hausse, embarquements en baisse vs mois précédent", "Les ventes accélèrent, la sortie physique traîne : congestion logistique, gros volumes à venir."),
        ("Vitória (ES)", "Port de l'Espírito Santo : surtout du conilon (robusta). Ligne à suivre pour l'ES."),
        ("REDEX / EADI Minas Gerais", "Entrepôts douaniers intérieurs du Minas. Certifié − dédouané = café vendu qui attend encore dans le Minas. Indicateur le plus direct de mouvement d'entrepôt au Minas."),
        ("Ne dit pas", "Le stock total des entrepôts du MG ou de l'ES : seule la part déjà vendue à l'export apparaît ici. Café des producteurs, coopératives non vendues et marché intérieur sont absents."),
        ("Méthode", "Jamais un jour isolé. Suivre chaque jour, comparer le cumul à la même période du mois précédent, puis au même mois de l'année précédente quand l'historique existera."),
    ]
    for a, b in reads:
        ws.cell(row=r, column=1, value=a).font = font(bold=True)
        ws.cell(row=r, column=2, value=b)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        for j in range(1, 4):
            ws.cell(row=r, column=j).alignment = Alignment(wrap_text=True, vertical="top"); ws.cell(row=r, column=j).border = BOX
        ws.row_dimensions[r].height = 40; r += 1
    r += 1
    ws.cell(row=r, column=1, value="Traduction du navigateur : « Victory » = Vitória, « Conillon » = conilon (robusta brésilien), « Redex and EADI » = entrepôts douaniers intérieurs.").font = font(italic=True, color="7A6F66")
    ws.cell(row=r + 1, column=1, value="Source : https://www.cecafe.com.br/dados-estatisticos/exportacoes-brasileiras/resumo-diario/").font = font(italic=True, color="7A6F66")
    ws.cell(row=r + 2, column=1, value="Onglets Arabica / Robusta : la cellule jaune « Date affichée » pilote tout ; par défaut le dernier jour saisi, tu peux y taper n'importe quelle date présente dans Donnees.").font = font(italic=True, color="7A6F66")
    widths(ws, [34, 70, 70])


def sheet_data(wb, rows):
    ws = wb.create_sheet("Donnees")
    cols = ["date", "table", "unit", "unit_label", "arabica_day", "conilon_day", "soluble_day", "total_day",
            "arabica_cum", "conilon_cum", "soluble_cum", "total_cum", "arabica_prev", "conilon_prev", "soluble_prev", "total_prev", "source"]
    header(ws, 1, cols)
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=dt.date.fromisoformat(r["date"])).number_format = "DD/MM/YYYY"
        ws.cell(row=i, column=2, value=r["table"]); ws.cell(row=i, column=3, value=r["unit"]); ws.cell(row=i, column=4, value=r["unit_label"])
        for j, k in enumerate(cols[4:16], start=5):
            c = ws.cell(row=i, column=j, value=int(float(r[k] or 0))); c.number_format = "#,##0"; c.font = BLUE
        ws.cell(row=i, column=17, value=r.get("source", "cecafe_daily.csv"))
        for j in range(1, 18): ws.cell(row=i, column=j).border = BOX
        if r["unit"] == "total":
            for j in range(1, 18): ws.cell(row=i, column=j).fill = SUB
    ws.freeze_panes = "E2"; widths(ws, [11, 12, 9, 30] + [10] * 12 + [34])
    ws.auto_filter.ref = f"A1:Q{max(2, len(rows) + 1)}"
    n = len(rows) + 3
    ws.cell(row=n, column=1, value="Colle ici les nouvelles journées (même format que volumes/cecafe_daily.csv). Les onglets Arabica et Robusta se mettent à jour seuls.").font = font(italic=True, color="7A6F66")


def combo(ws, anchor, cats, bar_refs, line_refs, ttl, ylab="sacs 60 kg"):
    bar = BarChart(); bar.type = "col"; bar.title = ttl; bar.y_axis.title = ylab; bar.y_axis.number_format = "#,##0"
    for ref, lab, rgb in bar_refs:
        bar.add_data(ref, titles_from_data=False); s = bar.series[-1]; s.tx = SeriesLabel(v=lab); color_bar(s, rgb)
    bar.set_categories(cats)
    if line_refs:
        line = LineChart()
        for ref, lab, rgb in line_refs:
            line.add_data(ref, titles_from_data=False); s = line.series[-1]; s.tx = SeriesLabel(v=lab); color_line(s, rgb)
        line.y_axis.axId = 200; line.y_axis.title = "variation vs mois précédent"; line.y_axis.number_format = "0%"
        line.y_axis.crosses = "max"; line.y_axis.majorGridlines = None
        bar += line
    bar.height = 9.5; bar.width = 20; bar.legend.position = "b"
    ws.add_chart(bar, anchor)


def sheet_type(wb, name, rows):
    dcol, ccol, pcol = TYPES[name]
    D = "Donnees"
    A = f"{D}!$A$2:$A${NMAX}"; B = f"{D}!$B$2:$B${NMAX}"; C = f"{D}!$C$2:$C${NMAX}"
    day = lambda date, stage, unit: f"SUMIFS({D}!${dcol}$2:${dcol}${NMAX},{A},{date},{B},\"{stage}\",{C},\"{unit}\")"
    cum = lambda date, stage, unit: f"SUMIFS({D}!${ccol}$2:${ccol}${NMAX},{A},{date},{B},\"{stage}\",{C},\"{unit}\")"
    prv = lambda date, stage, unit: f"SUMIFS({D}!${pcol}$2:${pcol}${NMAX},{A},{date},{B},\"{stage}\",{C},\"{unit}\")"
    ws = wb.create_sheet(name)
    kind = "arabica" if name == "Arabica" else "conilon (robusta)"
    title(ws, f"{name} — mouvement du jour, ports, cumul", f"Cecafé Resumo Diário · {kind} · sacs de 60 kg · tout est calculé à partir de l'onglet Donnees")
    ws["A4"] = "Date affichée"; ws["B4"] = f"=MAX({A})"; ws["B4"].fill = INP; ws["B4"].font = BLUE; ws["B4"].number_format = "DD/MM/YYYY"
    ws["C4"] = "← par défaut le dernier jour saisi ; tape une autre date (présente dans Donnees) pour la revoir"; ws["C4"].font = font(italic=True, color="7A6F66")
    ws["A5"] = "Veille (jour précédent saisi)"; ws["B5"] = f"=IF(MAX(INDEX(({A}<$B$4)*{A},0))=0,\"\",MAX(INDEX(({A}<$B$4)*{A},0)))"; ws["B5"].number_format = "DD/MM/YYYY"
    ws["A6"] = "Dernier jour saisi"; ws["B6"] = f"=MAX({A})"; ws["B6"].number_format = "DD/MM/YYYY"
    for c in ("A4", "A5", "A6"): ws[c].font = font(bold=True)
    DD, DV = "$B$4", "$B$5"

    # ---- block 0: simplified summary (movement of the day, current month, previous month, variation)
    rs = 8
    ws.cell(row=rs - 1, column=1, value=f"{name.upper()} — résumé (les noms de mois suivent la date affichée)").font = font(bold=True, size=11)
    header(ws, rs, ["Type", "Mouvement du jour", "", "", "Variation (%)"])
    ws.cell(row=rs, column=3, value=month_label(DD)); ws.cell(row=rs, column=4, value=month_label(DD, -1))
    for i, (k, lab) in enumerate([("certificates", "ÉMISSION (certificats)"), ("shipment", "EMBARQUE")]):
        r = rs + 1 + i
        ws.cell(row=r, column=1, value=lab).font = font(bold=True)
        ws.cell(row=r, column=2, value="=" + day(DD, k, "total")).number_format = "#,##0"
        ws.cell(row=r, column=3, value="=" + cum(DD, k, "total")).number_format = "#,##0"
        ws.cell(row=r, column=4, value="=" + prv(DD, k, "total")).number_format = "#,##0"
        ws.cell(row=r, column=5, value=f"=IF(D{r}=0,NA(),C{r}/D{r}-1)").number_format = "+0.0%;-0.0%;0.0%"
        ws.cell(row=r, column=5).font = font(bold=True)
        for j in range(1, 6): ws.cell(row=r, column=j).border = BOX; ws.cell(row=r, column=j).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
    ws.conditional_formatting.add(f"E{rs+1}:E{rs+2}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, bold=True, color="1E8449")))
    ws.conditional_formatting.add(f"E{rs+1}:E{rs+2}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, bold=True, color="C0392B")))
    ws.cell(row=rs + 3, column=1, value="Dédouanement = embarquement en total : même café vu par la douane. Colonne mois précédent = même période du mois précédent.").font = font(italic=True, color="7A6F66")

    # ---- block 1: day D vs prior day per stage, % vs previous month
    r0 = 13
    ws.cell(row=r0 - 1, column=1, value="1. Par étape, toutes unités : jour J, veille, cumul du mois, variation % vs mois précédent").font = font(bold=True, size=11)
    header(ws, r0, ["Étape", "Jour J", "Veille", "Cumul mois", "Mois précédent (même période)", "Variation % vs mois précédent", "Certifiés − embarqués (cumul)"])
    ws.cell(row=r0, column=4, value='="Cumul "&' + month_label(DD)[1:]); ws.cell(row=r0, column=5, value="=" + month_label(DD, -1)[1:] + '&" (même période)"')
    for i, (k, lab) in enumerate(STAGES):
        r = r0 + 1 + i
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=2, value="=" + day(DD, k, "total"))
        ws.cell(row=r, column=3, value=f"=IF({DV}=\"\",NA()," + day(DV, k, "total") + ")")
        ws.cell(row=r, column=4, value="=" + cum(DD, k, "total"))
        ws.cell(row=r, column=5, value="=" + prv(DD, k, "total"))
        ws.cell(row=r, column=6, value=f"=IF(E{r}=0,NA(),D{r}/E{r}-1)").number_format = "+0.0%;-0.0%;0.0%"
        for j in (2, 3, 4, 5): ws.cell(row=r, column=j).number_format = "#,##0"
        for j in range(1, 8): ws.cell(row=r, column=j).border = BOX
    ws.cell(row=r0 + 1, column=7, value=f"=D{r0+1}-D{r0+3}").number_format = "#,##0"
    ws.cell(row=r0 + 2, column=7, value="café vendu, pas encore parti").font = font(italic=True, color="7A6F66")
    cats = Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + 3)
    combo(ws, "I12", cats,
          [(Reference(ws, min_col=2, min_row=r0 + 1, max_row=r0 + 3), "Total du jour", RED),
           (Reference(ws, min_col=4, min_row=r0 + 1, max_row=r0 + 3), "Cumul du mois", GRN)],
          [(Reference(ws, min_col=6, min_row=r0 + 1, max_row=r0 + 3), "Variation % vs mois précédent", BLU)],
          f"{name} — mouvement du jour (rouge), cumul du mois (vert), variation % vs mois précédent (bleu, axe droit)")

    # ---- block 2: per port, day D per stage + % vs previous month
    r1 = r0 + 6
    ws.cell(row=r1 - 1, column=1, value="2 et 3. Par port : mouvement du jour (graphique 2) et variation % du cumul vs mois précédent (graphique 3)").font = font(bold=True, size=11)
    header(ws, r1, ["Port / unité", "Certificats J", "Dédouanés J", "Embarqués J", "Var. % certificats", "Var. % dédouanés", "Var. % embarqués"])
    for i, (k, lab) in enumerate(PORTS):
        r = r1 + 1 + i
        ws.cell(row=r, column=1, value=lab)
        for j, (st, _) in enumerate(STAGES):
            ws.cell(row=r, column=2 + j, value="=" + day(DD, st, k)).number_format = "#,##0"
            ws.cell(row=r, column=5 + j, value=f"=IF({prv(DD, st, k)}=0,NA(),{cum(DD, st, k)}/{prv(DD, st, k)}-1)").number_format = "+0.0%;-0.0%;0.0%"
        for j in range(1, 8): ws.cell(row=r, column=j).border = BOX
    ws.cell(row=r1 + len(PORTS) + 1, column=1, value="Vitória et REDEX/EADI Minas n'ont pas de ligne « embarquement » sur la page : leur café part par Santos. #N/A = pas de valeur le mois précédent.").font = font(italic=True, color="7A6F66")
    cats = Reference(ws, min_col=1, min_row=r1 + 1, max_row=r1 + len(PORTS))
    combo(ws, "I32", cats,
          [(Reference(ws, min_col=2, min_row=r1 + 1, max_row=r1 + len(PORTS)), "Certificats J", RED),
           (Reference(ws, min_col=3, min_row=r1 + 1, max_row=r1 + len(PORTS)), "Dédouanés J", RED2),
           (Reference(ws, min_col=4, min_row=r1 + 1, max_row=r1 + len(PORTS)), "Embarqués J", RED3)],
          [], f"{name} — mouvement du jour par port : certificats, dédouanés, embarqués")
    # chart 3: same ports × stages, but % change of the month-to-date vs previous month (bars up/down)
    pc = BarChart(); pc.type = "col"; pc.title = f"{name} — variation % du cumul vs mois précédent, par port"
    pc.y_axis.number_format = "0%"; pc.y_axis.title = "variation vs mois précédent"
    for j, (lab, rgb) in enumerate([("Certificats", BLU), ("Dédouanés", BLU2), ("Embarqués", BLU3)]):
        pc.add_data(Reference(ws, min_col=5 + j, min_row=r1 + 1, max_row=r1 + len(PORTS)), titles_from_data=False)
        sr = pc.series[-1]; sr.tx = SeriesLabel(v=lab); color_bar(sr, rgb)
    pc.set_categories(cats); pc.height = 9.5; pc.width = 20; pc.legend.position = "b"
    ws.add_chart(pc, "I52")

    # ---- block 3: per port, month-to-date vs previous month
    r2 = r1 + len(PORTS) + 4
    ws.cell(row=r2 - 1, column=1, value="4. Par port, cumul du mois (rouge) contre même période du mois précédent (bleu)").font = font(bold=True, size=11)
    header(ws, r2, ["Port / unité", "Certificats cumul", "Certificats mois préc.", "Dédouanés cumul", "Dédouanés mois préc.", "Embarqués cumul", "Embarqués mois préc."])
    for i, (k, lab) in enumerate(PORTS):
        r = r2 + 1 + i
        ws.cell(row=r, column=1, value=lab)
        for j, (st, _) in enumerate(STAGES):
            ws.cell(row=r, column=2 + 2 * j, value="=" + cum(DD, st, k)).number_format = "#,##0"
            ws.cell(row=r, column=3 + 2 * j, value="=" + prv(DD, st, k)).number_format = "#,##0"
        for j in range(1, 8): ws.cell(row=r, column=j).border = BOX
    cats = Reference(ws, min_col=1, min_row=r2 + 1, max_row=r2 + len(PORTS))
    combo(ws, "I72", cats,
          [(Reference(ws, min_col=2, min_row=r2 + 1, max_row=r2 + len(PORTS)), "Certificats cumul", RED),
           (Reference(ws, min_col=3, min_row=r2 + 1, max_row=r2 + len(PORTS)), "Certificats mois préc.", BLU),
           (Reference(ws, min_col=4, min_row=r2 + 1, max_row=r2 + len(PORTS)), "Dédouanés cumul", RED2),
           (Reference(ws, min_col=5, min_row=r2 + 1, max_row=r2 + len(PORTS)), "Dédouanés mois préc.", BLU2),
           (Reference(ws, min_col=6, min_row=r2 + 1, max_row=r2 + len(PORTS)), "Embarqués cumul", RED3),
           (Reference(ws, min_col=7, min_row=r2 + 1, max_row=r2 + len(PORTS)), "Embarqués mois préc.", BLU3)],
          [], f"{name} — cumul du mois (rouges) vs même période du mois précédent (bleus)")

    # ---- block 4: history, one row per day collected
    dates = sorted({r["date"] for r in rows})
    r3 = r2 + len(PORTS) + 4
    ws.cell(row=r3 - 1, column=1, value="5. Historique jour par jour (toutes unités)").font = font(bold=True, size=11)
    header(ws, r3, ["Date", "Certificats J", "Dédouanés J", "Embarqués J", "Certificats cumul", "Embarqués cumul", "Certifiés − embarqués", "Var. % certificats vs mois préc.", "Var. % embarqués vs mois préc."])
    for i, d in enumerate(dates):
        r = r3 + 1 + i
        ws.cell(row=r, column=1, value=dt.date.fromisoformat(d)).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=2, value="=" + day(f"$A{r}", "certificates", "total"))
        ws.cell(row=r, column=3, value="=" + day(f"$A{r}", "customs", "total"))
        ws.cell(row=r, column=4, value="=" + day(f"$A{r}", "shipment", "total"))
        ws.cell(row=r, column=5, value="=" + cum(f"$A{r}", "certificates", "total"))
        ws.cell(row=r, column=6, value="=" + cum(f"$A{r}", "shipment", "total"))
        ws.cell(row=r, column=7, value=f"=E{r}-F{r}")
        ws.cell(row=r, column=8, value=f"=IF({prv(f'$A{r}', 'certificates', 'total')}=0,NA(),E{r}/{prv(f'$A{r}', 'certificates', 'total')}-1)").number_format = "+0.0%;-0.0%;0.0%"
        ws.cell(row=r, column=9, value=f"=IF({prv(f'$A{r}', 'shipment', 'total')}=0,NA(),F{r}/{prv(f'$A{r}', 'shipment', 'total')}-1)").number_format = "+0.0%;-0.0%;0.0%"
        for j in (2, 3, 4, 5, 6, 7): ws.cell(row=r, column=j).number_format = "#,##0"
        for j in range(1, 10): ws.cell(row=r, column=j).border = BOX
    rl = r3 + max(1, len(dates))
    ws.cell(row=rl + 1, column=1, value="Le script build_xlsx.py ajoute une ligne par nouvelle journée. Si tu colles des jours à la main dans Donnees, ajoute aussi la date ici (colonne A), les formules de la ligne au-dessus se copient.").font = font(italic=True, color="7A6F66")
    cats = Reference(ws, min_col=1, min_row=r3 + 1, max_row=rl)
    combo(ws, "I92", cats,
          [(Reference(ws, min_col=2, min_row=r3 + 1, max_row=rl), "Certificats J", RED),
           (Reference(ws, min_col=3, min_row=r3 + 1, max_row=rl), "Dédouanés J", RED2),
           (Reference(ws, min_col=4, min_row=r3 + 1, max_row=rl), "Embarqués J", RED3)],
          [(Reference(ws, min_col=8, min_row=r3 + 1, max_row=rl), "Var. % certificats", BLU),
           (Reference(ws, min_col=9, min_row=r3 + 1, max_row=rl), "Var. % embarqués", BLU3)],
          f"{name} — historique quotidien (rouges) et variation % vs mois précédent (bleus, axe droit)")
    widths(ws, [30, 15, 15, 15, 20, 20, 20, 18, 18])


def main():
    rows = read_rows()
    wb = Workbook()
    sheet_definitions(wb)
    sheet_data(wb, rows)
    for name in ("Arabica", "Robusta"):
        sheet_type(wb, name, rows)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(OUT)
    print("wrote", os.path.relpath(OUT), "-", len(rows), "rows,", len({r['date'] for r in rows}), "day(s)")


if __name__ == "__main__":
    main()
